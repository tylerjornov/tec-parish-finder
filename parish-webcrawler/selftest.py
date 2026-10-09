#!/usr/bin/env python3
"""selftest.py - checks the tool's own logic. No internet, no Ollama, no website needed.

Run it with:   python selftest.py      (run.command can also run it for you)
If everything is fine you will see "ALL TESTS PASSED" at the end.

It tests: phone / email / address / person / URL normalizers, service-time sets (noon, ranges, "1st & 3rd"),
list comparison, status assignment (including UNCLEAR), missing markers, Asset Map contact
filtering, Rite derivation, the evidence "post-check" that keeps the AI honest, robots.txt wildcards,
Facebook wall detection, merging (rector vs church numbers, shared websites), and that the JSON round-trip
keeps unknown keys and key order. Also: the Id columns and results.json / results.csv / suggested_patch.json,
"N/A" and "None - No Rector", keeping the good part of a model answer, the "English" default, rite details and the
bishop table, church-name checks, addresses and cities, rector names without roles, shared diocese/college websites,
expired or taken-over websites, out-of-date pages, livestream links, skipping the model when the data is already on
the page, --fields / --ids-file, and the slowest-sites list in run.log.
"""

from __future__ import annotations

import json
import sys
import tempfile
import traceback
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import logging                                                                    # noqa: E402
logging.getLogger("parishcheck").addHandler(logging.NullHandler())                # keep the test output tidy
logging.getLogger("parishcheck").propagate = False

from parishcheck import models as M                                              # noqa: E402
from parishcheck.assetmap import listing_candidates, parse_listing               # noqa: E402
from parishcheck.cleaner import PageDoc, clean_html                              # noqa: E402
from parishcheck.comparators import compare_typed, decide_status                 # noqa: E402
from parishcheck.config import load_config                          # noqa: E402
from parishcheck.facebook import WALL_REASON, assess_recency, detect_wall, unwrap_fb_link  # noqa: E402
from parishcheck.llm import (                                                    # noqa: E402
    FieldAnswer, build_group_schema, clean_value, evidence_in_text, is_unclear, verify_value, _keyword_regex,
)
from parishcheck.merge import build_view, merge_field                            # noqa: E402
from parishcheck.models import Candidate                                         # noqa: E402
from parishcheck.normalizers import (                                            # noqa: E402
    addresses_match, derive_rite, extract_emails, extract_phones, format_phone, is_missing, livestream_key,
    norm_for_quote_check, normalize_url, parse_person, persons_match, phone_digits, role_is_parish_clergy,
    split_city_state, url_key,
)
from parishcheck.report import compare_record                                    # noqa: E402
from parishcheck.robots import RobotsRules                                       # noqa: E402
from parishcheck.rules import extract_rule_candidates, find_addresses, label_role  # noqa: E402
from parishcheck.times import (                                                  # noqa: E402
    find_times, format_time, parse_service_slots, restyle_times, time_keys,
)
from parishcheck.urls import classify_url, facebook_page_key, is_excluded, score_url  # noqa: E402

MARKERS = ["", "N/A - Data Not Available"]
RESULTS = {"passed": 0, "failed": 0}


def check(name: str, ok: bool, detail: str = "") -> None:
    if ok:
        RESULTS["passed"] += 1
    else:
        RESULTS["failed"] += 1
        print(f"  FAIL: {name}" + (f"   ({detail})" if detail else ""))


def status_of(ftype, json_value, site_values, **kw):
    return decide_status(ftype, json_value, site_values, missing_markers=MARKERS, **kw)


# --------------------------------------------------------------------------------------------------
def test_missing_markers():
    check("empty string is missing", is_missing("", MARKERS))
    check("None is missing", is_missing(None, MARKERS))
    check("marker is missing (any case, extra spaces)", is_missing("  n/a - data not available ", MARKERS))
    check("empty list is missing", is_missing([], MARKERS))
    check("real text is not missing", not is_missing("St. Andrew's", MARKERS))
    check("zero is a real value", not is_missing(0, MARKERS))
    check("custom marker works", is_missing("TBD", MARKERS + ["TBD"]))


def test_phones():
    check("phone formats", phone_digits("(864) 235-5884") == "8642355884" == phone_digits("+1 864.235.5884"))
    check("phone extension ignored", phone_digits("864-235-5884 ext. 12") == "8642355884")
    check("not a phone", phone_digits("12345") is None)
    check("find phones in text", extract_phones("Call (864) 235-5884 or 864-555-1212; fax 800 555 0000") ==
          ["8642355884", "8645551212", "8005550000"])
    check("zip+4 / dates are not phones", extract_phones("Greenville SC 29601-1234, posted 2025-10-08") == [])
    check("format_phone", format_phone("8642355884") == "(864) 235-5884")
    check("phone equal", compare_typed("phone", "864-235-5884", ["(864) 235-5884"]).status == M.CORRECT)
    check("phone differs", compare_typed("phone", "864-235-5884", ["(864) 235-9999"]).status == M.DISCREPANCY)
    check("phone any of several", compare_typed("phone", "864-111-2222", ["(864) 235-5884", "(864) 111-2222"]).status == M.CORRECT)
    check("json holds two phones", compare_typed("phone", "864-111-2222 / 864-235-5884", ["(864) 235-5884"]).status == M.CORRECT)


def test_emails():
    check("email case-insensitive", compare_typed("email", "Office@Church.ORG", ["office@church.org"]).status == M.CORRECT)
    check("email any candidate", compare_typed("email", "b@x.org", ["a@x.org", "b@x.org"]).status == M.CORRECT)
    check("email differs", compare_typed("email", "b@x.org", ["a@x.org"]).status == M.DISCREPANCY)
    check("junk emails dropped", extract_emails("logo@2x.png sentry@sentry.io real@church.org") == ["real@church.org"])


def test_urls():
    check("url ignores scheme/www/slash", url_key("https://www.Church.org/about/") == url_key("http://church.org/about"))
    check("tracking params stripped", normalize_url("https://x.org/a?utm_source=z&fbclid=1#frag") == "https://x.org/a")
    check("index.html folded", normalize_url("https://x.org/index.html") == "https://x.org/")
    check("relative links", normalize_url("../a", "https://x.org/b/c") == "https://x.org/a")
    check("mailto is not a page", normalize_url("mailto:a@b.org") == "")
    check("website: host only", compare_typed("url", "http://church.org/", ["https://www.church.org/home"], url_kind="website").status == M.CORRECT)
    check("website: different host", compare_typed("url", "http://old.org/", ["https://new.org/"], url_kind="website").status == M.DISCREPANCY)
    check("livestream: youtube channel forms", livestream_key("https://www.youtube.com/@abc/live") == livestream_key("youtube.com/@abc"))
    check("livestream matches among links", compare_typed("url", "https://www.youtube.com/@abc/live",
          ["https://vimeo.com/1", "https://www.youtube.com/@abc"], url_kind="livestream").status == M.CORRECT)
    check("classify facebook", classify_url("https://m.facebook.com/stjohns") == "facebook")
    check("classify asset map", classify_url("https://www.episcopalassetmap.org/dioceses/x/list/y") == "asset_map")
    check("classify own", classify_url("stjohns.org") == "own_site")
    check("classify missing", classify_url("  ") == "missing")
    check("classify unsupported", classify_url("https://instagram.com/x", ["instagram.com"]) == "unsupported")
    check("fb page key", facebook_page_key("https://www.facebook.com/StJohnsChurch/about") == "stjohnschurch")
    check("fb profile id", facebook_page_key("https://www.facebook.com/profile.php?id=1000123") == "id:1000123")
    check("fb share links unsupported", facebook_page_key("https://www.facebook.com/sharer/sharer.php?u=x") == "")
    rx = [__import__("re").compile(p, __import__("re").I) for p in [r"/(?:19|20)\d{2}/(?:0?[1-9]|1[0-2])(?:/|$)", r"/(?:tag|category)(?:/|$)", r"/wp-json"]]
    check("exclude date archive", is_excluded("https://x.org/2024/03/post", rx))
    check("exclude tag pages", is_excluded("https://x.org/tag/music", rx))
    check("exclude calendar query", is_excluded("https://x.org/events?tribe-bar-date=2025-01", rx))
    check("keep normal page", not is_excluded("https://x.org/worship", rx))
    check("priority score", score_url("https://x.org/worship", "", ["worship"]) > score_url("https://x.org/blog/a/b/c", "", ["worship"]))


def test_addresses():
    check("St/Street, S/South, punctuation", addresses_match("1002 S Main St., Greenville, SC", "1002 South Main Street, Greenville, SC 29601"))
    check("Saint Charles", addresses_match("123 St Charles Ave, New Orleans, LA", "123 Saint Charles Avenue, New Orleans, Louisiana"))
    check("different number", not addresses_match("1002 S Main St", "1004 S Main St"))
    check("different direction", not addresses_match("1002 S Main St., Greenville, SC", "1002 N Main St, Greenville, SC"))
    check("different city", not addresses_match("10 Main St, Austin, TX", "10 Main St, Dallas, TX"))
    check("city missing on one side is fine", addresses_match("10 Main St", "10 Main St, Dallas, TX"))
    check("Rd/Road, Ave/Avenue, Dr/Drive, Hwy/Highway", all([
        addresses_match("5 Elm Rd", "5 Elm Road"), addresses_match("5 Oak Ave", "5 Oak Avenue"),
        addresses_match("5 Pine Dr", "5 Pine Drive"), addresses_match("5 US Hwy 1", "5 US Highway 1")]))
    check("address compare typed", compare_typed("address", "1002 S Main St., Greenville, SC", ["22 Elm St, X, SC", "1002 South Main Street, Greenville, SC"]).status == M.CORRECT)
    check("city/state split", split_city_state("1002 S Main St., Greenville, SC 29601") == ("Greenville", "SC"))
    found = find_addresses("Visit us at 1002 S. Main Street, Greenville, SC 29601 or write to PO Box 5.")
    check("find address in text", len(found) == 1 and found[0][0] == "1002 S. Main Street, Greenville, SC", str(found))
    found2 = find_addresses("Parish Hall\n2803 1st St\nWyandotte, MI 48192")
    check("find address across lines", found2 and found2[0][0] == "2803 1st St, Wyandotte, MI", str(found2))


def test_times_and_services():
    ts = lambda s: sorted(format_time(t.minutes) for t in find_times(s))
    check("noon / midnight", ts("12:00 noon and midnight") == ["12:00 midnight", "12:00 noon"], str(ts("12:00 noon and midnight")))
    check("a.m./p.m. forms", ts("8 a.m., 5:30 pm, 10.30 AM") == ["10:30 AM", "5:30 PM", "8:00 AM"], str(ts("8 a.m., 5:30 pm, 10.30 AM")))
    check("range keeps start", [format_time(s.minutes) for s in parse_service_slots("9-10:30 am Forum", "sun")] == ["9:00 AM"])
    check("range crossing noon", ts("11:00-1:00 PM") == ["11:00 AM", "1:00 PM"], str(ts("11:00-1:00 PM")))
    check("12:30 AM / PM", ts("12:30 AM 12:30 PM") == ["12:30 AM", "12:30 PM"])
    check("bare 12-hour guess", ts("Mass at 5:30 and 8:00") == ["5:30 PM", "8:00 AM"], str(ts("Mass at 5:30 and 8:00")))
    check("24h output", restyle_times("Mass 5:30 PM, noon, 8:30am", "24h") == "Mass 17:30, 12:00, 08:30", restyle_times("Mass 5:30 PM, noon, 8:30am", "24h"))
    check("12h output", restyle_times("at 8:30am and 7pm", "12h") == "at 8:30 AM and 7:00 PM")
    check("John 3:16 left alone", restyle_times("John 3:16 study", "12h", only_explicit=True) == "John 3:16 study")
    check("Midnight Mass title not a time", ts("Midnight Mass 11:30 PM") == ["11:30 PM"])
    gold_sun = "8:30 AM Low Mass, Rite II; 10:30 AM Sung High Mass with Incense, Rite I (livestreamed)"
    gold_wk = "Wed 5:30 PM Healing Mass (Benediction of the Blessed Sacrament following, every 4th Wednesday); 3rd Wed 5:00 PM Rosary"
    desc = lambda text, day=None: sorted(__import__("parishcheck.times", fromlist=["describe_slot"]).describe_slot(s) for s in parse_service_slots(text, day))
    check("sunday default day", desc(gold_sun, "sun") == ["Sun 10:30 AM", "Sun 8:30 AM"], str(desc(gold_sun, "sun")))
    check("parenthetical ordinal ignored; 3rd Wed kept", desc(gold_wk) == ["3rd Wed 5:00 PM", "Wed 5:30 PM"], str(desc(gold_wk)))
    check("1st & 3rd", desc("1st & 3rd Thursdays at 10 a.m.") == ["1st Thu 10:00 AM", "3rd Thu 10:00 AM"], str(desc("1st & 3rd Thursdays at 10 a.m.")))
    check("Mon-Fri range", len(parse_service_slots("Mon–Fri 12:00 noon Holy Eucharist")) == 5)
    check("two days two times", desc("Tue 7:00 AM Morning Prayer, Wed 5:30 PM Eucharist") == ["Tue 7:00 AM", "Wed 5:30 PM"])
    # comparisons
    check("services equal", compare_typed("services", gold_sun, ["10:30 AM Sung High Mass, Rite I; 8:30 AM Low Mass, Rite II"], default_day="sun").status == M.CORRECT)
    r = compare_typed("services", gold_sun, ["8:30 AM Low Mass; 5:00 PM Evensong"], default_day="sun")
    check("services overlap = PARTIAL", r.status == M.PARTIAL and r.site_only and r.json_only, str(r))
    check("services none in common = DISCREPANCY", compare_typed("services", "9:00 AM Eucharist", ["11:00 AM Eucharist"], default_day="sun").status == M.DISCREPANCY)
    check("same time, different week-of-month = PARTIAL", compare_typed("services", "Wed 5:30 PM Mass", ["1st Wed 5:30 PM Mass"]).status == M.PARTIAL)
    check("services with no clock times fall back to text", compare_typed("services", "Sunday Eucharist", ["Sunday Eucharist"]).status == M.CORRECT)


def test_lists_person_rite():
    check("list correct (fuzzy)", compare_typed("list", "English; Spanish", ["Spanish", "english"]).status == M.CORRECT)
    r = compare_typed("list", "English; Spanish", ["English"])
    check("list partial", r.status == M.PARTIAL and r.json_only == ["Spanish"], str(r))
    check("list extra on site is partial", compare_typed("list", "English", ["English", "Spanish"]).status == M.PARTIAL)
    check("list discrepancy", compare_typed("list", "Latin", ["English"]).status == M.DISCREPANCY)
    check("text fuzzy match", compare_typed("text", "Parking lot on site", ["On-site parking lot"]).status == M.CORRECT)
    check("person ignores honorifics", persons_match("The Reverend J. Gary Eichelberger, Rector", "Fr. Gary Eichelberger")[0])
    check("person: a different role is still the same person", persons_match("The Rev. John Bishop, Vicar", "Rev. John Bishop, Rector")[0])
    check("person different surname", not persons_match("The Rev. Jane Smith, Rector", "The Rev. Jane Jones, Rector")[0])
    check("person same surname, different first name", not persons_match("The Rev. Jane Smith, Rector", "The Rev. John Smith, Rector")[0])
    p = parse_person("The Reverend J. Gary Eichelberger, Rector")
    check("person parse", p.surname == "eichelberger" and p.roles == frozenset({"rector"}), str(p))
    check("interim rector is its own role", parse_person("Rev. A. Lee, Interim Rector").roles == frozenset({"interim rector"}))
    check("surname Bishop survives", parse_person("The Rev. John Bishop, Vicar").surname == "bishop")
    check("rite I", derive_rite("8:00 Rite I") == "I")
    check("rite II", derive_rite("10:30 Rite Two") == "II")
    check("rite mixed", derive_rite("Rite I", "Rite II") == "Mixed")
    check("rite I/II", derive_rite("Rite I/II") == "Mixed")
    check("rite none", derive_rite("Holy Eucharist") is None)
    check("rite II not mistaken for I", derive_rite("Rite II only") == "II")
    check("rite compare", compare_typed("rite", "Mixed", ["Mixed"]).status == M.CORRECT)


def test_status_assignment():
    s = status_of
    check("CORRECT", s("phone", "864-235-5884", ["(864) 235-5884"]).status == M.CORRECT)
    check("DISCREPANCY", s("phone", "864-235-5884", ["(864) 999-9999"]).status == M.DISCREPANCY)
    check("BOTH_MISSING (empty + marker)", s("text", "N/A - Data Not Available", []).status == M.BOTH_MISSING)
    check("BOTH_MISSING None", s("text", None, []).status == M.BOTH_MISSING)
    check("SITE_ONLY when JSON holds a marker", s("text", "N/A - Data Not Available", ["Parking lot"]).status == M.SITE_ONLY)
    check("JSON_ONLY", s("text", "Parking lot", [], json_only_note="crawl hit cap").status == M.JSON_ONLY)
    check("JSON_ONLY note kept", "cap" in s("text", "x", [], json_only_note="crawl hit cap").detail)
    u = s("text", "x", [], unclear_reason="blocked by Facebook login wall")
    check("UNCLEAR when blocked", u.status == M.UNCLEAR and "login wall" in u.reason)
    check("UNCLEAR beats BOTH_MISSING when source failed", s("text", "", [], unclear_reason="fetch failed").status == M.UNCLEAR)
    low = s("text", "Parking lot", ["Street parking"], site_confidence="low")
    check("low confidence difference -> UNCLEAR", low.status == M.UNCLEAR and "low confidence" in low.reason, str(low))
    check("low confidence agreement stays CORRECT", s("text", "Parking lot", ["Parking lot"], site_confidence="low").status == M.CORRECT)
    check("low confidence SITE_ONLY -> UNCLEAR", s("text", "", ["Street parking"], site_confidence="low").status == M.UNCLEAR)
    p = s("list", "English; Spanish", ["English"])
    check("PARTIAL lists the items", p.status == M.PARTIAL and p.json_only == ["Spanish"])
    check("state compare", compare_typed("state", "South Carolina", ["SC"]).status == M.CORRECT)


def test_asset_map():
    check("Rector counts", role_is_parish_clergy("Rector"))
    check("Vicar counts", role_is_parish_clergy("Vicar"))
    check("Priest-in-Charge counts", role_is_parish_clergy("Priest-in-Charge"))
    check("Interim Rector counts", role_is_parish_clergy("Interim Rector"))
    check("Dean counts", role_is_parish_clergy("Dean"))
    check("Canon to the Ordinary does NOT count", not role_is_parish_clergy("Canon to the Ordinary"))
    check("Parish Administrator does NOT count", not role_is_parish_clergy("Parish Administrator"))
    check("Assistant to the Rector does NOT count", not role_is_parish_clergy("Assistant to the Rector"))
    check("Associate Rector does NOT count", not role_is_parish_clergy("Associate Rector"))
    check("empty title does NOT count", not role_is_parish_clergy(""))
    html = """<html><body><main><h1>St. Test Church</h1>
    <p class="address"><span class="address-line1">1 Main St</span><br><span class="locality">Testville</span>, <span class="administrative-area">MI</span> <span class="postal-code">48000</span></p>
    <div class="field--name-field-phone"><a href="tel:734-555-0100">734-555-0100</a></div>
    <div class="field--name-field-email-address"><a href="mailto:office@test.org">office@test.org</a></div>
    <div class="field--name-field-external-url"><a href="http://www.test.org">http://www.test.org</a></div>
    <div class="block-field_contacts" data-block-plugin-id="field_block:node:place:field_contacts"><h2>Contact</h2>
      <div class="details--person"><h3><div class="field--name-field-person-name">Canon Pat Diocese</div></h3><div class="field--name-field-person-position">Canon to the Ordinary</div><div class="details--person__email"><a href="mailto:canon@diocese.org">canon@diocese.org</a></div></div>
      <div class="details--person"><h3><div class="field--name-field-person-name">Pastor Lee Priest</div></h3><div class="field--name-field-person-position">Rector</div><div class="details--person__email"><a href="mailto:lee@test.org">lee@test.org</a></div></div>
    </div>
    <div class="worship-times"><div class="field--name-field-worship-times-day">Sunday</div><div class="field--name-field-worship-times-time">10:00 am</div></div>
    </main></body></html>"""
    listing = parse_listing(html, "https://www.episcopalassetmap.org/x")
    check("AM address parsed", listing.address == "1 Main St, Testville, MI", listing.address)
    check("AM website parsed", listing.website == "http://www.test.org/", listing.website)
    cands, notes = listing_candidates(listing, listing.url)
    clergy = [c for c in cands if c.field == "@clergy"]
    check("AM: only the Rector becomes rector", len(clergy) == 1 and "Lee Priest" in clergy[0].value, str([c.value for c in clergy]))
    rector_emails = [c.value for c in cands if c.field == "@email" and c.role == "rector"]
    check("AM: canon's email is never a rector email", rector_emails == ["lee@test.org"], str(rector_emails))
    check("AM: ignored contact is explained", any("Canon to the Ordinary" in n for n in notes), str(notes))
    check("AM: church email stays church", any(c.field == "@email" and c.role == "church" and c.value == "office@test.org" for c in cands))


def test_postcheck():
    page = "Sunday Worship: 8:00 AM Holy Eucharist Rite I. The Rev. Jane Doe, Rector. Call (864) 235-5884 or office@church.org."
    pn = norm_for_quote_check(page)
    check("evidence found (case/punctuation ignored)", evidence_in_text("the rev. jane doe, RECTOR", pn))
    check("evidence with ... pieces", evidence_in_text("Sunday Worship ... Holy Eucharist Rite I", pn))
    check("evidence invented", not evidence_in_text("Sunday 11:00 AM Mass", pn))
    check("empty evidence rejected", not evidence_in_text("", pn))
    from parishcheck.config import FieldSpec
    svc = FieldSpec("sunday_services", "services")
    check("time present", verify_value(svc, "8:00 AM Holy Eucharist, Rite I", page, pn)[0])
    check("time invented is rejected", not verify_value(svc, "9:30 AM Choral Eucharist", page, pn)[0])
    check("noon / 12h loose match", verify_value(svc, "8:00 AM Eucharist", "Eucharist at 8 a.m.", norm_for_quote_check("Eucharist at 8 a.m."))[0])
    per = FieldSpec("rector_name", "person")
    check("name present", verify_value(per, "The Reverend Jane Doe, Rector", page, pn)[0])
    check("name invented is rejected", not verify_value(per, "The Reverend John Smith, Rector", page, pn)[0])
    check("email invented is rejected", not verify_value(FieldSpec("e", "text"), "write to boss@church.org", page, pn)[0])
    check("phone present", verify_value(FieldSpec("p", "text"), "Call (864) 235-5884", page, pn)[0])
    check("phone invented is rejected", not verify_value(FieldSpec("p", "text"), "Call (864) 555-1111", page, pn)[0])
    check("unclear words", all(is_unclear(x) for x in ["unclear", "", "N/A", "Unknown", "not stated"]))
    check("real answer is not unclear", not is_unclear("Parking lot on site"))
    cfg = tiny_config()
    rector = FieldSpec("rector_name", "person", group="clergy")
    check("rector in house style", clean_value(rector, "Reverend Jane Doe, Rector", cfg) == "The Rev. Jane Doe",
          clean_value(rector, "Reverend Jane Doe, Rector", cfg))
    check("dean keeps Very Rev. and Dr.", clean_value(rector, "The Very Reverend Dr. Jane Doe, Dean", cfg) == "The Very Rev. Dr. Jane Doe")
    clergy = FieldSpec("other_clergy", "list", group="clergy")
    got = clean_value(clergy, "The Reverend Ann Lee, Associate Rector; Father Bob Ray (Curate)", cfg)
    check("other clergy keep roles in parentheses", got == "The Rev. Ann Lee (Associate Rector); Fr. Bob Ray (Curate)", got)
    ans = FieldAnswer.model_validate({"lines": [3, 4], "value": "y", "confidence": "high"})
    check("pydantic answer", ans.confidence == "high" and ans.lines == [3, 4])
    schema = build_group_schema([FieldSpec("music_style", "text"), FieldSpec("parking", "text")])
    check("schema asks for lines, value and confidence for every field",
          schema["required"] == ["music_style", "parking"] and
          list(schema["properties"]["parking"]["properties"]) == ["lines", "value", "confidence"], str(schema))

    from parishcheck.digest import build_digest, page_lines
    rx = _keyword_regex(["parking", "re:\\d{1,2}\\s?[ap]\\.?m\\b"])
    menu = "Home\nAbout\nFree parking behind the church"
    a = PageDoc("https://church.org/visit", title="Visit", text="Welcome\nWe are glad you came\nService at 10 am\nCoffee after\n" + "x\n" * 20 + menu)
    b = PageDoc("https://church.org/about", title="About", text="Our history\nFounded long ago\n" + menu)
    dg = build_digest([(0, a), (5, b)], rx, 3500, 1)
    texts = [ln.text for ln in dg.lines]
    check("digest keeps keyword lines with one line of context", texts[:3] == ["We are glad you came", "Service at 10 am", "Coffee after"], str(texts))
    check("digest shows a repeated menu/footer line only once", texts.count("Free parking behind the church") == 1 and len(dg.pages) == 1, str(texts))
    check("digest numbers lines and labels pages", dg.text.startswith("=== Page P1: church.org/visit") and "2: Service at 10 am" in dg.text, dg.text)
    check("digest neighbours stay on the cited page", [ln.n for ln in dg.neighbours(2)] == [1, 2, 3])
    check("digest respects the budget", sum(len(t) for t in [ln.text for ln in build_digest([(0, a)], rx, 30, 1).lines]) <= 30)
    check("digest: no keyword -> nothing", not build_digest([(0, b)], _keyword_regex(["choir"]), 3500, 1).lines)
    check("long paragraphs are cut into short lines", all(len(x) <= 281 for x in page_lines(PageDoc("u", text="Word. " * 200))))


def test_robots_and_facebook():
    r = RobotsRules("User-agent: *\nDisallow: /maps\nDisallow: /*/map\nDisallow: /node/*/edit\nAllow: /maps/public$\nCrawl-delay: 4\nSitemap: https://x.org/sm.xml\n")
    check("robots disallow prefix", not r.allowed("/maps/anything"))
    check("robots wildcard", not r.allowed("/en/map") and not r.allowed("/node/12/edit"))
    check("robots allow ordinary", r.allowed("/worship"))
    check("robots longest match wins", r.allowed("/maps/public"))
    check("robots crawl-delay and sitemap", r.crawl_delay == 4 and r.sitemaps == ["https://x.org/sm.xml"])
    named = RobotsRules("User-agent: ParishInfoVerifier\nDisallow: /\n\nUser-agent: *\nAllow: /\n")
    check("robots group that names us wins", not named.allowed("/anything"))
    check("empty robots allows all", RobotsRules("").allowed("/x"))
    real_page = ("Our Church About\n" + "We are a welcoming parish with many programs and a long history in the town. " * 8)
    check("fb: real content is not a wall", detect_wall("https://www.facebook.com/x/about", 200, real_page) == "")
    check("fb: login redirect is a wall", detect_wall("https://www.facebook.com/login/?next=x", 200, real_page) == WALL_REASON)
    check("fb: captcha text is a wall", detect_wall("https://www.facebook.com/x", 200, real_page + " Confirm you're human") == WALL_REASON)
    check("fb: only a login box is a wall", detect_wall("https://www.facebook.com/x", 200, "Log in\nEmail or phone\nPassword\nCreate new account") == WALL_REASON)
    check("fb: 403 is a wall", detect_wall("https://www.facebook.com/x", 403, real_page) == WALL_REASON)
    check("fb: l.php unwrap", unwrap_fb_link("https://l.facebook.com/l.php?u=https%3A%2F%2Fchurch.org%2F&h=AT0") == "https://church.org/")
    today = date(2026, 10, 8)
    check("fb: recent relative date", assess_recency("Sunday service at 10! 3d", today) == "recent")
    check("fb: old full date", assess_recency("March 3, 2024 Easter service at 9", today) == "stale")
    check("fb: new full date", assess_recency("September 14, 2026 Bring a friend", today) == "recent")
    check("fb: years-ago relative", assess_recency("Pancake supper 2y", today) == "stale")
    check("fb: no clues", assess_recency("Join us Sunday", today) == "unknown")


def test_sitemaps():
    import gzip
    from parishcheck.crawler import _parse_sitemap
    from parishcheck.fetcher import _maybe_gunzip
    index = b"<?xml version='1.0'?><sitemapindex xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'><sitemap><loc>https://x.org/page-sitemap.xml.gz</loc></sitemap><sitemap><loc>https://x.org/post-sitemap.xml</loc></sitemap></sitemapindex>"
    urls = b"<?xml version='1.0'?><urlset xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'><url><loc>https://x.org/</loc></url><url><loc> https://x.org/worship </loc></url></urlset>"
    kind, locs = _parse_sitemap(index)
    check("nested sitemap index is recognised", kind == "index" and locs == ["https://x.org/page-sitemap.xml.gz", "https://x.org/post-sitemap.xml"], str((kind, locs)))
    kind, locs = _parse_sitemap(urls)
    check("sitemap urls are read (spaces trimmed)", kind == "urls" and locs == ["https://x.org/", "https://x.org/worship"], str(locs))
    check(".gz sitemap is unzipped", _maybe_gunzip(gzip.compress(urls)) == urls)
    check("plain data is left alone", _maybe_gunzip(urls) == urls)
    check("corrupt .gz does not crash (returns nothing)", _maybe_gunzip(b"\x1f\x8bnot really gzip") == b"")
    check("broken XML does not crash", _parse_sitemap(b"<urlset><url><loc>https://x.org/a</loc></url")[1] == ["https://x.org/a"])
    check("garbage does not crash", _parse_sitemap(b"hello") == ("urls", []))


def test_placeholder_sites_and_json_ld():
    from parishcheck.crawler import dead_site_reason
    dead = clean_html("<html><body><h1>Account Suspended</h1><p>This account has been suspended.</p></body></html>", "http://x.org/")
    check("suspended-hosting page is not a real website", bool(dead_site_reason(dead, "http://x.org/")))
    check("tempdisabled address is not a real website", bool(dead_site_reason(dead, "http://x.org/cgi-sys/tempdisabled.cgi")))
    real = clean_html("<html><body><h1>Welcome</h1><p>" + "We are a friendly parish with many programs. " * 40 + "</p></body></html>", "http://x.org/")
    check("a normal website is accepted", dead_site_reason(real, "http://x.org/") == "")
    html = ('<html><head><script type="application/ld+json">{"@type":"Church","name":"X","address":"Greenville, SC 29601 USA"}</script></head>'
            '<body><main>' + "Plenty of text about our church and its programs. " * 20 + '</main></body></html>')
    cands = extract_rule_candidates(clean_html(html, "https://x.org/"), [])
    check("JSON-LD address without a street is ignored", not any(c.field == "@address" for c in cands), str([c.value for c in cands if c.field == "@address"]))
    html2 = html.replace('"address":"Greenville, SC 29601 USA"', '"address":{"@type":"PostalAddress","streetAddress":"1002 S Main St.","addressLocality":"Greenville","addressRegion":"SC"}')
    cands2 = extract_rule_candidates(clean_html(html2, "https://x.org/"), [])
    check("JSON-LD PostalAddress is used", any(c.field == "@address" and c.value == "1002 S Main St., Greenville, SC" for c in cands2), str([c.value for c in cands2]))


def test_json_ld_wrappers():
    from parishcheck.cleaner import _jsonld_blocks, _parse
    body = '{"@type": "Church", "name": "X"}'
    want = [{"@type": "Church", "name": "X"}]
    for label, raw in [("<!-- -->", f"<!-- {body} -->"), ("<!-- --!>", f"<!-- {body} --!>"),
                       ("CDATA", f"//<![CDATA[ {body} //]]>"), ("plain", body)]:
        got = _jsonld_blocks(_parse(f'<html><head><script type="application/ld+json">{raw}</script></head></html>'))
        check(f"JSON-LD wrapped in {label} parses", got == want, str(got))


def test_cleaner_rules_merge():
    html = """<html><head><title>Contact</title></head><body><main><h1>Contact us</h1>
    <p>The Rev. Jane Doe, Rector<br>Phone: (864) 235-5890<br><a href="mailto:rector@church.org">rector@church.org</a></p>
    <p>Parish Office: (864) 235-5884 <a href="mailto:office@church.org">office@church.org</a> Fax: (864) 235-0000</p>
    <p>Watch live: <a href="https://www.youtube.com/@church/live">Livestream</a>. Find us at 12 Elm Street, Greenville, SC 29601.
    Plenty of extra words so that the extractor keeps this paragraph in the main text of the page for sure.</p>
    </main><footer>12 Elm Street<br>Greenville, SC 29601<br>(864) 235-5884</footer></body></html>"""
    doc = clean_html(html, "https://church.org/contact", ["youtube.com"])
    cands = extract_rule_candidates(doc, ["youtube.com"])
    roles = {(c.field, c.value, c.role) for c in cands}
    check("rector phone labelled", ("@phone", "(864) 235-5890", "rector") in roles, str(roles))
    check("office phone labelled church", ("@phone", "(864) 235-5884", "church") in roles)
    check("fax number skipped", not any("0000" in c.value for c in cands))
    check("rector email labelled", ("@email", "rector@church.org", "rector") in roles)
    check("office email labelled church", ("@email", "office@church.org", "church") in roles)
    check("livestream found", any(c.field == "@livestream" and "youtube.com/@church/live" in c.value for c in cands))
    check("address found", any(c.field == "@address" and c.value == "12 Elm Street, Greenville, SC" for c in cands))
    cfg = tiny_config()
    church_phone = merge_field(cfg.get_field("church_phone"), cands, cfg)
    rector_phone = merge_field(cfg.get_field("rector_phone"), cands, cfg)
    check("church_phone picks the office line", church_phone.best == "(864) 235-5884", str(church_phone))
    check("rector_phone picks the rector line", rector_phone.best == "(864) 235-5890", str(rector_phone))
    church_email = merge_field(cfg.get_field("church_email"), cands, cfg)
    check("church_email never the rector's", church_email.best == "office@church.org", str(church_email))
    # rector_phone JSON that is really the office number: found on site but labelled for the other role
    sv_rector_email = merge_field(cfg.get_field("rector_email"), cands, cfg)
    check("rector_email", sv_rector_email.best == "rector@church.org")


def tiny_config():
    yaml_text = """
output_dir: out_test
id_key: id
name_key: name
website_key: website
uncheckable_fields: [id, lat, notes, date_last_updated]
missing_markers: ["", "N/A - Data Not Available", "N/A"]
confirmed_none_markers: ["None - No Rector"]
diocese_lookup:
  Upper South Carolina: The Rt. Rev. Daniel P. Richards
fields:
  - {key: name, type: text, group: identity, description: "name"}
  - {key: website, type: url, group: contact, url_kind: website, description: "site"}
  - {key: address, type: address, group: contact, description: "address"}
  - {key: city, type: text, group: contact, address_part: city, description: "city"}
  - {key: state, type: text, group: contact, address_part: state, description: "state"}
  - {key: church_phone, type: phone, group: contact, role: church, description: "phone"}
  - {key: church_email, type: email, group: contact, role: church, description: "email"}
  - {key: rector_email, type: email, group: contact, role: rector, description: "rector email"}
  - {key: rector_phone, type: phone, group: contact, role: rector, description: "rector phone"}
  - {key: rector_name, type: person, group: clergy, role: rector, description: "rector"}
  - {key: sunday_services, type: services, group: services, default_day: sun, description: "sunday"}
  - {key: weekday_services, type: services, group: services, description: "weekday", confirmed_none: ["N/A"]}
  - {key: rite_details, type: rite, group: services, description: "rite details"}
  - {key: rite, type: rite, group: services, description: "rite"}
  - {key: service_languages, type: list, group: services, description: "languages", default_value: English, default_unless: [spanish, español, bilingual, korean]}
  - {key: diocesan_bishop, type: person, group: lookup, lookup: diocese, description: "bishop"}
  - {key: music_style, type: text, group: community, description: "music"}
  - {key: other_contact, type: list, group: derived, description: "other"}
derived_fields:
  other_contact: {from: [church_phone, church_email], separator: " | "}
  rite_details: {details_from_services: [sunday_services, weekday_services]}
  rite: {from_services: [sunday_services, weekday_services]}
"""
    d = Path(tempfile.mkdtemp(prefix="parish_selftest_"))
    (d / "config.yaml").write_text(yaml_text, encoding="utf-8")
    return load_config(d / "config.yaml")


def cand(field, value, role=None, page="https://church.org/", rank=0, conf="high", kind=M.OWN_SITE, method="rule", evidence="ev"):
    return Candidate(field=field, value=value, role=role, source_url=page, page_rank=rank, confidence=conf,
                     source_type=kind, method=method, evidence=evidence)


def result(kind, cands, status="ok", reason="", url="https://church.org/", **kw):
    r = {"key": "k", "kind": kind, "url": url, "status": status, "reason": reason, "pages_crawled": 3, "cap_hit": False,
         "candidates": [c.to_dict() for c in cands], "issues": [], "notes": []}
    r.update(kw)
    return r


def test_merge_and_rows():
    cfg = tiny_config()
    rec = {"id": "a1", "name": "St. A", "website": "https://church.org", "address": "12 Elm St, Greenville, SC", "city": "Greenville",
           "church_phone": "864-235-5884", "church_email": "N/A - Data Not Available", "rector_email": "rector@church.org",
           "sunday_services": "8:00 AM Eucharist", "music_style": "Choir and organ", "notes": "keep me"}
    cands = [
        cand("@phone", "(864) 235-5884", "church"), cand("@email", "office@church.org", "church"),
        cand("@address", "12 Elm Street, Greenville, SC"), cand("@website", "https://church.org/"),
        cand("sunday_services", "8:00 AM Holy Eucharist, Rite I; 10:30 AM Choral Eucharist, Rite II", page="https://church.org/worship"),
        cand("music_style", "Gospel band", conf="low", method="llm"),
    ]
    view = build_view(rec, cfg, [result(M.OWN_SITE, cands)])
    rows = {r.field: r for r in compare_record(rec, cfg, view, "a1", "St. A")}
    check("phone CORRECT", rows["church_phone"].status == M.CORRECT)
    check("email SITE_ONLY", rows["church_email"].status == M.SITE_ONLY and rows["church_email"].site_value == "office@church.org")
    check("address CORRECT", rows["address"].status == M.CORRECT)
    check("city taken from the address", rows["city"].status == M.CORRECT, rows["city"].status)
    check("services PARTIAL", rows["sunday_services"].status == M.PARTIAL and "10:30" in rows["sunday_services"].detail, rows["sunday_services"].detail)
    check("low-confidence difference is UNCLEAR", rows["music_style"].status == M.UNCLEAR and rows["music_style"].reason)
    check("uncheckable fields get no row", "id" not in rows and "notes" not in rows and "lat" not in rows, str(sorted(rows)))
    check("rows carry the JSON city", rows["church_phone"].city == "Greenville", rows["church_phone"].city)
    check("rector_phone not on site and not in JSON", rows["rector_phone"].status == M.BOTH_MISSING)
    check("rite derived from services text", rows["rite"].status == M.SITE_ONLY and rows["rite"].site_value == "Mixed", rows["rite"].site_value)
    check("other_contact is never 'missing from your data'", rows["other_contact"].status == M.BOTH_MISSING, rows["other_contact"].status)
    check("row carries method/source", rows["church_phone"].method == "rule" and rows["church_phone"].source_type == "own site")

    # failed source -> everything UNCLEAR with a reason
    bad = build_view(rec, cfg, [result(M.FACEBOOK, [], status="blocked", reason="blocked by Facebook login wall")])
    brows = compare_record(rec, cfg, bad, "a1", "St. A")
    check("blocked Facebook -> all checkable fields UNCLEAR", all(r.status == M.UNCLEAR and "login wall" in r.reason
          for r in brows if r.field not in ("other_contact", "diocesan_bishop")), str([(r.field, r.status) for r in brows if r.status != M.UNCLEAR]))

    # source priority: own site beats Facebook; disagreement is mentioned in detail
    own = result(M.OWN_SITE, [cand("@phone", "(864) 235-5884", "church")])
    fb = result(M.FACEBOOK, [cand("@phone", "(864) 111-0000", "church", kind=M.FACEBOOK)], url="https://facebook.com/x")
    v2 = build_view(rec, cfg, [fb, own])
    check("own site wins over Facebook", v2.values["church_phone"].best == "(864) 235-5884")
    check("Facebook disagreement noted", any("Facebook says" in n for n in v2.values["church_phone"].notes), str(v2.values["church_phone"].notes))

    # JSON_ONLY notes when the page cap was hit
    cap = result(M.OWN_SITE, [], cap_hit=True)
    v3 = build_view(rec, cfg, [cap])
    r3 = {r.field: r for r in compare_record(rec, cfg, v3, "a1", "St. A")}
    check("JSON_ONLY mentions page cap", r3["church_phone"].status == M.JSON_ONLY and "page limit" in r3["church_phone"].detail, r3["church_phone"].detail)

    # shared website: choose the location page that shows THIS record's address
    rec_b = dict(rec, id="b2", address="99 Oak Road, Springfield, SC", church_phone="803-000-1111")
    shared_c = [
        cand("@address", "12 Elm Street, Greenville, SC", page="https://church.org/greenville"),
        cand("@phone", "(864) 235-5884", "church", page="https://church.org/greenville"),
        cand("@address", "99 Oak Road, Springfield, SC", page="https://church.org/springfield"),
        cand("@phone", "(803) 000-1111", "church", page="https://church.org/springfield"),
    ]
    vb = build_view(rec_b, cfg, [result(M.OWN_SITE, shared_c)], shared=True)
    check("shared site: second location gets ITS phone", vb.values["church_phone"].best == "(803) 000-1111", str(vb.values["church_phone"]))
    check("shared site noted", any("shared website: multiple locations" in n for n in vb.notes))
    rec_c = dict(rec, id="c3", address="1 Nowhere Lane, Elsewhere, SC", church_phone="864-777-8888")
    vc = build_view(rec_c, cfg, [result(M.OWN_SITE, shared_c)], shared=True)
    rc = {r.field: r for r in compare_record(rec_c, cfg, vc, "c3", "St. C")}
    check("shared site, location unknown: differences become UNCLEAR", rc["church_phone"].status == M.UNCLEAR, rc["church_phone"].status)

    # conflicting rector answers -> UNCLEAR
    two = [cand("rector_name", "The Rev. Ann Lee, Rector", method="llm"), cand("rector_name", "The Rev. Bob Ray, Rector", method="llm")]
    vv = build_view(rec, cfg, [result(M.OWN_SITE, two)])
    rr = {r.field: r for r in compare_record(dict(rec, rector_name="The Rev. Cy Fox, Rector"), cfg, vv, "a1", "St. A")}
    check("conflicting people -> UNCLEAR", rr["rector_name"].status == M.UNCLEAR and "different answers" in rr["rector_name"].reason, rr["rector_name"].reason)

    # rejected model answers -> UNCLEAR
    rj = result(M.OWN_SITE, [], issues=[["parking", "rejected_unverified", "quote not found", "https://church.org/x"]])
    v4 = build_view(rec, cfg, [rj])
    check("issues recorded", "parking" in v4.issues)


def _row(field, status, site_value, conf):
    from parishcheck.report import Row
    return Row("a1", "St. A", field, status, "", site_value, confidence=conf)


class FakeModel:
    """Stands in for Ollama: returns canned answers so the post-check can be tested without any AI."""

    def __init__(self, *answers):
        self.answers = list(answers)
        self.calls = 0

    def chat_json(self, system, user, schema, num_predict=700):
        self.calls += 1
        if not self.answers:
            return None, "no more answers"
        a = self.answers.pop(0) if len(self.answers) > 1 or self.calls == 1 else self.answers[0]
        return a, ("" if a is not None else "the model's answer was not valid JSON")


def ask(model, group_fields, page_text, group="community", cfg=None):
    """Runs the post-check on one page. The test answers quote their evidence; here each quote is turned into
    the number of the digest line that contains it (or a line number that does not exist)."""
    import re as _re
    from parishcheck.digest import build_digest
    from parishcheck.llm import get_group_info, run_group_on_digest
    cfg = cfg or tiny_config()
    cfg.style_examples = {"music_style": ["Traditional Anglican choral music, with a professional organist"],
                          "accessibility": ["Handicapped-accessible facilities"]}
    cfg.groups = {"community": {"keywords": ["music", "choir", "organ", "accessib", "wheelchair", "service", "worship"]}}
    specs = [f for f in cfg.fields if f.key in group_fields]
    info = get_group_info(cfg, group, specs)
    lined = _re.sub(r"(?<=\.) ", "\n", page_text)
    doc = PageDoc("https://church.org/x", title="T", text=lined, main_text=lined)
    dg = build_digest([(0, doc)], _re.compile("."), 3500, 1)

    def to_lines(a):
        if not isinstance(a, dict):
            return a
        out = {}
        for k, v in a.items():
            ev = norm_for_quote_check(v.get("evidence", ""))
            nums = [ln.n for ln in dg.lines if ev and ev in norm_for_quote_check(ln.text)] or ([99] if ev else [])
            out[k] = {"lines": nums, "value": v["value"], "confidence": v["confidence"]}
        return out
    model.answers = [to_lines(a) for a in model.answers]
    return run_group_on_digest(model, cfg, info, specs, dg, source_type=M.OWN_SITE, today=date(2026, 10, 8))


def ans(evidence, value, conf="high"):
    return {"evidence": evidence, "value": value, "confidence": conf}


def test_model_answers_are_checked():
    cfg = tiny_config()
    cfg.fields.append(__import__("parishcheck.config", fromlist=["FieldSpec"]).FieldSpec("accessibility", "text", group="community"))
    page = ("Worship at St. A. Sunday worship: 8:00 AM Holy Eucharist Rite I. Our choir sings hymns each week with the organ. "
            "The building is wheelchair accessible. The Rev. Jane Doe, Rector. Mother Jane leads the staff.")
    good = {"music_style": ans("Our choir sings hymns each week with the organ", "Choir singing hymns, with organ"),
            "accessibility": ans("The building is wheelchair accessible", "Handicapped-accessible facilities")}
    out = ask(FakeModel(good), ["music_style", "accessibility"], page, cfg=cfg)
    got = {c.field: c.value for c in out.candidates}
    check("good model answers are accepted", got.get("music_style") == "Choir singing hymns, with organ" and
          got.get("accessibility") == "Handicapped-accessible facilities", str(got) + str(out.issues))
    copied = {"music_style": ans("Worship at St. A.", "Traditional Anglican choral music, with a professional organist"),
              "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(copied), ["music_style", "accessibility"], page, cfg=cfg)
    check("copied style example + unrelated quote is rejected", not out.candidates and any(i[0] == "music_style" for i in out.issues), str(out.issues))
    fake_quote = {"music_style": ans("The choir performs Bach cantatas every Sunday", "Choir performs Bach cantatas"), "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(fake_quote), ["music_style", "accessibility"], page, cfg=cfg)
    check("invented quote is rejected (rejected_unverified)", not out.candidates and out.issues and out.issues[0][1] == "rejected_unverified", str(out.issues))
    low = {"music_style": ans("", "Some music", "low"), "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(low), ["music_style", "accessibility"], page, cfg=cfg)
    check("low-confidence unverifiable answer is quietly 'unclear'", not out.candidates and not out.issues, str(out.issues))
    out = ask(FakeModel(None, None), ["music_style", "accessibility"], page, cfg=cfg)
    check("malformed JSON twice -> model_failed", not out.candidates and out.issues and all(i[1] == "model_failed" for i in out.issues), str(out.issues))
    flaky = FakeModel(None, good)
    out = ask(flaky, ["music_style", "accessibility"], page, cfg=cfg)
    check("malformed JSON once -> retried and recovered", len(out.candidates) == 2 and flaky.calls == 2, f"calls={flaky.calls} {out.issues}")

    # services: one invented item must not throw away the good ones
    sv = {"sunday_services": ans("Sunday worship: 8:00 AM Holy Eucharist Rite I", "8:00 AM Holy Eucharist, Rite I; 11:00 AM Sung High Mass with Incense, Rite II"),
          "weekday_services": ans("", "unclear", "low"), "rite_details": ans("", "unclear", "low"), "service_languages": ans("", "unclear", "low")}
    out = ask(FakeModel(sv), ["sunday_services", "weekday_services", "rite_details", "service_languages"], page, group="community", cfg=cfg)
    got = {c.field: c.value for c in out.candidates}
    check("services: invented item dropped, real item kept", got.get("sunday_services") == "8:00 AM Holy Eucharist, Rite I", str(got) + str(out.issues))

    # a verified answer the model itself was unsure about is kept, but marked low (so it becomes UNCLEAR)
    unsure = {"music_style": ans("Our choir sings hymns each week with the organ", "Choir with organ", "low"), "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(unsure), ["music_style", "accessibility"], page, cfg=cfg)
    check("low-confidence but verified answer is kept as low", out.candidates and out.candidates[0].confidence == "low", str(out.issues))
    # evidence has to be about the topic of the group
    off_topic = {"music_style": ans("The Rev. Jane Doe, Rector", "Choir with organ"), "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(off_topic), ["music_style", "accessibility"], page, cfg=cfg)
    check("evidence about a different topic is rejected", not out.candidates and out.issues, str(out.issues))


def test_model_breaks_phone_ties():
    from parishcheck.extract import ExtractContext, disambiguate_leftovers
    cfg = tiny_config()
    cands = [cand("@phone", "(864) 111-0001", None, page="https://church.org/a", evidence="Call (864) 111-0001"),
             cand("@phone", "(864) 111-0002", None, page="https://church.org/b", evidence="Call (864) 111-0002")]
    fake = FakeModel({"church_main": 1, "rector": 0})
    ctx = ExtractContext(cfg=cfg, client=fake, state=None, today=date(2026, 10, 8), stop=__import__("threading").Event())
    extra = disambiguate_leftovers(cands, ctx, M.OWN_SITE)
    roles = {(c.value, c.role, c.method, c.confidence) for c in extra}
    check("model labels look-alike phones (low confidence, method llm)",
          roles == {("(864) 111-0002", "church", "llm", "low"), ("(864) 111-0001", "rector", "llm", "low")}, str(roles))
    labelled = [cand("@phone", "(864) 111-0001", "church"), cand("@phone", "(864) 111-0002", None)]
    check("no model call when a church number is already labelled", disambiguate_leftovers(labelled, ctx, M.OWN_SITE) == [])
    sv = merge_field(cfg.get_field("church_phone"), cands + extra, cfg)
    check("a model-only label makes the merged value low confidence", sv.best == "(864) 111-0002" and sv.confidence == "low", str(sv))


def test_files_and_encoding():
    from parishcheck.pipeline import load_records
    from parishcheck.config import ConfigError
    d = Path(tempfile.mkdtemp(prefix="parish_selftest_enc_"))
    cfg = tiny_config()
    cfg.input_json = d / "p.json"
    cfg.input_json.write_bytes(("\ufeff" + json.dumps([{"id": "1", "name": "Iglesia San José – Español", "x": "naïve café"}], ensure_ascii=False)).encode("utf-8"))
    recs = load_records(cfg)
    check("JSON with a BOM and accents loads", recs[0]["name"] == "Iglesia San José – Español" and recs[0]["x"] == "naïve café")
    cfg.input_json.write_text("{not json", encoding="utf-8")
    try:
        load_records(cfg)
        check("bad JSON gives a friendly error", False)
    except ConfigError as exc:
        check("bad JSON gives a friendly error", "not valid JSON" in str(exc), str(exc))
    cfg.input_json.write_text('{"a": 1}', encoding="utf-8")
    try:
        load_records(cfg)
        check("non-list JSON gives a friendly error", False)
    except ConfigError as exc:
        check("non-list JSON gives a friendly error", "list" in str(exc))
    cfg.input_json = d / "missing.json"
    try:
        load_records(cfg)
        check("missing file gives a friendly error", False)
    except ConfigError as exc:
        check("missing file gives a friendly error", "can't find" in str(exc))
    from parishcheck.report import write_excel, write_results_files, Row
    from openpyxl import load_workbook
    cfg.output_dir = d / "out"
    rows = [Row("1", "Iglesia San José", "name", M.DISCREPANCY, "a", "b", source_url="https://sanjose.org/about",
                method="llm", confidence="high", evidence="Iglesia b", city="Greenville", verified=True),
            Row("1", "Iglesia San José", "city", M.SITE_ONLY, "", "=Greenville\x01", method="rule (from address)",
                confidence="high", evidence="x" * 5000, city="Greenville"),
            Row("2", "A Church", "name", M.CORRECT, "a", "a", city="Aiken"),
            Row("0", "A Church", "rector_name", M.SITE_ONLY, "", "The Rev. Ann Lee", method="llm", confidence="high", city="Camden"),
            Row("0", "A Church", "parking", M.SITE_ONLY, "", "Lot", method="llm", confidence="low", city="Camden")]
    summary = [{"id": "1", "name": "Iglesia San José", "website": "https://sanjose.org", "pages_crawled": 3, "cap_hit": False,
                "counts": {M.DISCREPANCY: 1, M.SITE_ONLY: 1}, "notes": "", "pdfs": ["https://sanjose.org/bulletin.pdf"],
                "site_problem": ""}]
    meta = {"run_at": "2026-10-09T10:00:00-04:00", "input_file": "p.json", "model": "qwen2.5:3b", "parishes": 3,
            "fields": [], "ids_file": "", "only": ""}
    path = write_excel(cfg, rows, summary, meta)
    check("only the workbook in the output folder", [p.name for p in cfg.output_dir.iterdir() if not p.name.startswith(".")] == [path.name])
    wb = load_workbook(path)
    check("workbook sheets", wb.sheetnames == ["To review", "Summary", "All checks", "Key"], str(wb.sheetnames))
    rv = wb["To review"]
    check("To review has only review rows, accents intact", rv.max_row == 5 and rv["A2"].value == "A Church", str(rv.max_row))
    check("To review: Id next to Parish, sorted by name then id",
          [c.value for c in rv[1]][:4] == ["Parish", "Id", "Field", "Problem"] and [rv["B2"].value, rv["B4"].value] == ["0", "1"],
          str([[c.value for c in r][:2] for r in rv.iter_rows()]))
    check("To review has Confidence and Evidence", [c.value for c in rv[1]][8:10] == ["Confidence", "Evidence (words on the page)"]
          and rv["I4"].value == "high" and rv["J4"].value == "Iglesia b", str([c.value for c in rv[1]]))
    check("plain-English problem label", rv["D4"].value == "Different on website", str(rv["D4"].value))
    check("source is a clickable link", rv["H4"].hyperlink is not None and rv["H4"].hyperlink.target == "https://sanjose.org/about")
    check("site text starting with = is not a formula", rv["F5"].value == "=Greenville" and rv["F5"].data_type == "s", repr(rv["F5"].value))
    check("All checks has every row and an Id column", wb["All checks"].max_row == 6 and wb["All checks"]["B1"].value == "Id")
    sm = wb["Summary"]
    check("summary starts with the run, file and model", "qwen2.5:3b" in sm["A1"].value and "p.json" in sm["A1"].value, sm["A1"].value)
    check("summary has Id and Site problem columns", [c.value for c in sm[2]][:4] == ["Parish", "Id", "Website", "Site problem"])
    check("summary lists PDFs", "bulletin.pdf" in (sm["I3"].value or ""))

    paths = write_results_files(cfg, rows, meta)
    check("results.json, results.csv and suggested_patch.json are written",
          sorted(p.name for p in paths) == ["results.csv", "results.json", "suggested_patch.json"])
    res = json.loads((cfg.output_dir / "results.json").read_text(encoding="utf-8"))
    check("results.json has meta (run, file, model)", res["meta"]["model"] == "qwen2.5:3b" and res["meta"]["input_file"] == "p.json"
          and res["meta"]["run_at"].startswith("2026"), str(res["meta"]))
    first = res["results"][0]
    check("results.json keys", list(first) == ["id", "name", "city", "field", "status", "label", "json_value", "site_value",
                                               "notes", "source_url", "method", "confidence", "evidence"], str(list(first)))
    check("results.json sorted by name then id", [(r["name"], r["id"]) for r in res["results"]][:3] ==
          [("A Church", "0"), ("A Church", "0"), ("A Church", "2")], str([(r["name"], r["id"]) for r in res["results"]]))
    check("results.json keeps accents and is not cut short", any(r["name"] == "Iglesia San José" and len(r["evidence"]) == 5000
          for r in res["results"]))
    import csv as _csv
    with open(cfg.output_dir / "results.csv", encoding="utf-8", newline="") as fh:
        csv_rows = list(_csv.DictReader(fh))
    check("results.csv has the same rows", len(csv_rows) == 5 and csv_rows[0]["status"] == "SITE_ONLY" and csv_rows[0]["label"] == "Missing from your data")
    patch = json.loads((cfg.output_dir / "suggested_patch.json").read_text(encoding="utf-8"))
    got = sorted((p["id"], p["field"]) for p in patch)
    check("patch: only high-confidence SITE_ONLY/DISCREPANCY found by a rule or verified",
          got == [("1", "city"), ("1", "name")], str(got))
    check("patch format", set(patch[0]) == {"id", "field", "old", "new", "source_url", "evidence"}, str(patch[0]))


def test_record_selection():
    from parishcheck.pipeline import select_indices
    cfg = tiny_config()
    recs = [{"id": "st-andrews-greenville", "name": "St. Andrew's", "website": "https://a.org"},
            {"id": "grace-1", "name": "Grace Church", "website": "https://b.org"},
            {"id": "grace-2", "name": "Grace Two", "website": "https://b.org/x"},
            {"id": "ch-3", "name": "Christ", "website": "https://www.facebook.com/christchurch"}]
    check("--only matches id or name substring (any case)", select_indices(recs, cfg, "ANDREWS") == [0] and select_indices(recs, cfg, "grace") == [1, 2])
    check("no --only selects everything", select_indices(recs, cfg, None) == [0, 1, 2, 3])


# --------------------------------------------------------------------------------------------------
# Fewer false alarms (P2)
# --------------------------------------------------------------------------------------------------
def _rows_for(rec, cands, cfg=None, **kw):
    cfg = cfg or tiny_config()
    view = build_view(rec, cfg, [result(M.OWN_SITE, cands, **kw)])
    return {r.field: r for r in compare_record(rec, cfg, view, rec.get("id", "a1"), rec.get("name", "St. A"))}


BASE = {"id": "a1", "name": "St. A", "website": "https://church.org", "address": "12 Elm St, Greenville, SC", "city": "Greenville"}


def test_other_contact_is_never_missing():
    real = load_config(HERE / "config.yaml")
    check("other_contact is not checked in config.yaml", real.get_field("other_contact") is None)
    cands = [cand("@phone", "(864) 235-5884", "church"), cand("@email", "office@church.org", "church")]
    rec = dict(BASE, church_phone="864-235-5884", church_email="office@church.org")
    r = _rows_for(rec, cands)["other_contact"]
    check("other_contact empty -> not 'missing from your data'", r.status == M.BOTH_MISSING, r.status)
    r = _rows_for(dict(rec, other_contact="(864) 235-5884 | office@church.org"), cands)["other_contact"]
    check("other_contact compared with the JSON's own phone and email", r.status == M.CORRECT and r.method.startswith("rule"), r.status)
    r = _rows_for(dict(rec, other_contact="(803) 000-0000"), cands)["other_contact"]
    check("other_contact that disagrees with the JSON's own phone", r.status == M.DISCREPANCY, r.status)


def test_na_and_confirmed_none():
    cfg = tiny_config()
    check("'N/A' counts as missing (any case)", is_missing("N/A", cfg.missing_markers) and is_missing(" n/a ", cfg.missing_markers))
    nobody = dict(BASE, rector_name="None - No Rector")
    r = _rows_for(nobody, [])["rector_name"]
    check("'None - No Rector' and no rector on the site -> Matches", r.status == M.CORRECT, f"{r.status} {r.detail}")
    r = _rows_for(nobody, [cand("rector_name", "The Rev. Ann Lee", method="llm")])["rector_name"]
    check("'None - No Rector' but the site names one -> Different, with a note", r.status == M.DISCREPANCY
          and "JSON says confirmed none" in r.detail, f"{r.status} {r.detail}")
    sunday_only = dict(BASE, weekday_services="N/A")
    r = _rows_for(sunday_only, [cand("sunday_services", "10:00 AM Holy Eucharist", method="llm")])["weekday_services"]
    check("weekday 'N/A' and no weekday services on the site -> Matches", r.status == M.CORRECT, f"{r.status} {r.detail}")
    r = _rows_for(sunday_only, [cand("weekday_services", "Wed 12:00 noon Holy Eucharist", method="llm")])["weekday_services"]
    check("weekday 'N/A' but the site lists one -> Different", r.status == M.DISCREPANCY and "confirmed none" in r.detail, r.status)
    r = _rows_for(dict(BASE, parking="N/A"), [])
    check("'N/A' elsewhere is just missing", "parking" not in r or r["parking"].status == M.BOTH_MISSING)


def test_verifier_keeps_the_good_part():
    from parishcheck.llm import drop_rites, salvage
    page = "Sunday worship: 8:00 AM Holy Eucharist Rite I. Our choir sings hymns each week with the organ."
    pn = norm_for_quote_check(page)
    from parishcheck.config import FieldSpec
    svc = FieldSpec("sunday_services", "services")
    chk = lambda t: verify_value(svc, t, page, pn)
    check("a failing (parenthetical) is dropped, the rest kept",
          salvage("8:00 AM Holy Eucharist (livestream at 9:30 AM)", chk, with_commas=False)[0] == "8:00 AM Holy Eucharist")
    check("a failing core still fails", salvage("9:30 AM Choral Eucharist (spoken)", chk, with_commas=False)[0] == "")
    check("services never shrink to a bare day", salvage("Mon, Wed 9:30 AM Eucharist", chk, with_commas=False)[0] == "")
    check("unsupported Rite is dropped", drop_rites("8:30 AM Low Mass, Rite II", {"II"}) == "8:30 AM Low Mass")
    check("supported Rite is kept", drop_rites("8:30 AM Low Mass, Rite I", {"II"}) == "8:30 AM Low Mass, Rite I")
    cfg = tiny_config()
    sv = {"sunday_services": ans("Sunday worship: 8:00 AM Holy Eucharist Rite I", "8:00 AM Holy Eucharist, Rite II (with incense)"),
          "weekday_services": ans("", "unclear", "low"), "service_languages": ans("", "unclear", "low")}
    out = ask(FakeModel(sv), ["sunday_services", "weekday_services", "service_languages"], page, cfg=cfg)
    got = {c.field: c.value for c in out.candidates}
    check("model answer: wrong Rite and invented detail dropped, service kept", got.get("sunday_services") == "8:00 AM Holy Eucharist",
          str(got) + str(out.issues))
    good = {"music_style": ans("Our choir sings hymns each week with the organ", "Choir sings hymns, with harpsichord and trumpets"),
            "accessibility": ans("", "unclear", "low")}
    cfg.fields.append(__import__("parishcheck.config", fromlist=["FieldSpec"]).FieldSpec("accessibility", "text", group="community"))
    out = ask(FakeModel(good), ["music_style", "accessibility"], page, cfg=cfg)
    check("list/text: an invented extra after a comma is dropped", [c.value for c in out.candidates] == ["Choir sings hymns"],
          str([c.value for c in out.candidates]) + str(out.issues))
    check("verified answers are marked for the patch", out.candidates and out.candidates[0].extra.get("verified") is True)
    copied = {"music_style": ans("Our choir sings hymns each week with the organ", "Traditional Anglican choral music, with a professional organist"),
              "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(copied), ["music_style", "accessibility"], page, cfg=cfg)
    check("a long style example is never accepted as an answer", all("Traditional" not in c.value for c in out.candidates),
          str([c.value for c in out.candidates]))


class _NoState:
    def llm_get(self, key):
        return None

    def llm_put(self, key, value):
        pass


def _run_groups(cfg, model, text):
    from parishcheck.extract import ExtractContext, new_result, run_llm_groups
    ctx = ExtractContext(cfg=cfg, client=model, state=_NoState(), today=date(2026, 10, 8), stop=__import__("threading").Event())
    doc = PageDoc("https://church.org/worship", title="Worship", text=text, main_text=text)
    res = new_result(M.OWN_SITE, "k", "https://church.org/")
    return run_llm_groups([doc], M.OWN_SITE, ctx, ["services"], res), res


class PromptModel(FakeModel):
    def chat_json(self, system, user, schema, num_predict=700):
        self.prompts = getattr(self, "prompts", []) + [user]
        return super().chat_json(system, user, schema, num_predict)


def test_languages_default_and_rite_details():
    cfg = tiny_config()
    cfg.groups = {"services": {"keywords": ["sunday", "eucharist", "service", "spanish"]}}
    model = PromptModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"}})
    cands, res = _run_groups(cfg, model, "Sunday service\n10:00 AM Holy Eucharist\nAll are welcome")
    eng = [c for c in cands if c.field == "service_languages"]
    check("no other language on the site -> 'English' by rule, low confidence",
          len(eng) == 1 and eng[0].value == "English" and eng[0].method == "rule (default)" and eng[0].confidence == "low", str(eng))
    check("... and the model is not asked about languages", model.calls == 1 and "service_languages" not in model.prompts[0])
    model = PromptModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "service_languages": {"lines": [2], "value": "English; Spanish", "confidence": "high"}})
    cands, _ = _run_groups(cfg, model, "Sunday service\n1:00 PM Holy Eucharist in Spanish\nAll are welcome")
    check("another language on the site -> the model is asked", "service_languages" in model.prompts[0])
    check("'English' needs not be on the page", [c.value for c in cands if c.field == "service_languages"] == ["English; Spanish"],
          str([(c.field, c.value) for c in cands]))
    default = cand("service_languages", "English", method="rule (default)", conf="low")
    r = _rows_for(dict(BASE), [default])["service_languages"]
    check("assumed 'English' and blank JSON -> not a review row", r.status == M.BOTH_MISSING and r.site_value == "English", r.status)
    r = _rows_for(dict(BASE, service_languages="English"), [default])["service_languages"]
    check("assumed 'English' agrees with JSON 'English'", r.status == M.CORRECT, r.status)
    svc = cand("sunday_services", "8:00 AM Holy Eucharist, Rite I; 10:30 AM Choral Eucharist, Rite II", method="llm")
    rows = _rows_for(dict(BASE, rite_details="Rite I and Rite II both offered"), [svc])
    check("rite_details worked out from the services text", rows["rite_details"].site_value == "Rite I (8:00 AM); Rite II (10:30 AM)"
          and rows["rite_details"].method.startswith("rule"), rows["rite_details"].site_value)
    check("rite_details compared by which rites are used", rows["rite_details"].status == M.CORRECT, rows["rite_details"].status)
    check("rite_details is never asked of the model", "rite_details" not in [f.key for g in
          __import__("parishcheck.extract", fromlist=["x"]).llm_specs_by_group(cfg).values() for f in g])


def test_diocese_lookup():
    rec = dict(BASE, diocese="Upper South Carolina", diocesan_bishop="The Rt. Rev. Daniel P. Richards")
    r = _rows_for(rec, [])["diocesan_bishop"]
    check("bishop matches the lookup table", r.status == M.CORRECT and r.method == "rule (diocese lookup)", r.status)
    r = _rows_for(dict(rec, diocesan_bishop="The Rt. Rev. Andrew Waldo"), [])["diocesan_bishop"]
    check("old bishop -> Different on website", r.status == M.DISCREPANCY and r.site_value == "The Rt. Rev. Daniel P. Richards")
    r = _rows_for(dict(rec, diocese="Atlantis"), [])["diocesan_bishop"]
    check("unknown diocese -> Couldn't tell, saying why", r.status == M.UNCLEAR and "diocese_lookup" in r.reason)
    r = _rows_for(rec, [], status="failed", reason="the website could not be opened")["diocesan_bishop"]
    check("the lookup works even when the website fails", r.status == M.CORRECT, r.status)
    check("the bishop is never asked of the model", "diocesan_bishop" not in [f.key for g in
          __import__("parishcheck.extract", fromlist=["x"]).llm_specs_by_group(tiny_config()).values() for f in g])


def test_church_names():
    from parishcheck.normalizers import looks_like_church_name
    check("web designer credit is not a name", not looks_like_church_name("Digital Pros"))
    check("a code is not a name", not looks_like_church_name("PG168K"))
    check("'The Episcopal Church' / a diocese is not a parish name", not looks_like_church_name("The Episcopal Church")
          and not looks_like_church_name("The Episcopal Diocese of South Carolina"))
    check("real names pass", all(looks_like_church_name(n) for n in ["St. Anne's Conway", "Christ Church", "Grace Church Anderson",
          "Church of the Holy Communion", "Trinity Cathedral"]))
    v = merge_field(tiny_config().get_field("name"), [cand("name", "Digital Pros", method="llm"), cand("@jsonld_name", "PG168K")], tiny_config())
    check("junk names are dropped before merging", v is None, str(v))
    rec = dict(BASE, name="Grace Episcopal Church")
    check("name: dropping 'Episcopal' only is a match", compare_typed("name", "St. Anne's Episcopal Church", ["St. Anne's Church"]).status == M.CORRECT)
    check("name: word order only is a match", compare_typed("name", "Episcopal Church of the Epiphany", ["Church of the Epiphany, Episcopal"]).status == M.CORRECT)
    r = _rows_for(rec, [cand("name", "Grace Church Anderson", method="llm")])["name"]
    check("name: same name worded differently -> Partly matches", r.status == M.PARTIAL, r.status)
    check("name: a different saint is Different", compare_typed("name", "St. Anne's Church", ["St. Mark's Church"]).status == M.DISCREPANCY)
    check("name: Saint = St., apostrophes ignored", compare_typed("name", "Saint Philips Church", ["St. Philip's Episcopal Church"]).status == M.CORRECT)


def test_addresses_and_cities():
    check("Twelfth = 12th", addresses_match("1001 12th St, Cayce, SC", "1001 Twelfth Street, Cayce, SC"))
    check("Twenty-First = 21st", addresses_match("100 21st Ave", "100 Twenty-First Avenue"))
    check("Street = St", addresses_match("3001 Meeting Street", "3001 Meeting St."))
    check("Hilton Head = Hilton Head Island", addresses_match("3001 Meeting St, Hilton Head, SC", "3001 Meeting Street, Hilton Head Island, SC"))
    check("a different town is still different", not addresses_match("10 Main St, Columbia, SC", "10 Main St, Cayce, SC"))
    glued = find_addresses("Office: (843) 681-8333 3001 Meeting St, Hilton Head Island, SC 29926")
    check("a phone number is not glued onto the street number", [a for a, _ in glued] == ["3001 Meeting St, Hilton Head Island, SC"], str(glued))
    glued2 = find_addresses("Call 843 681 8333 3001 Meeting St., Hilton Head Island, SC")
    check("... also with spaces in the phone number", [a for a, _ in glued2] == ["3001 Meeting St., Hilton Head Island, SC"], str(glued2))
    rec = dict(BASE, city="Greenville")
    cands = [cand("@address", "1 Office Rd, Columbia, SC", page="https://church.org/a"), cand("@address", "1 Office Rd, Columbia, SC", page="https://church.org/b"),
             cand("@address", "12 Elm Street, Greenville, SC", page="https://church.org/c")]
    r = _rows_for(rec, cands)["city"]
    check("city: the JSON city among the site's addresses -> Matches", r.status == M.CORRECT, f"{r.status} {r.detail}")
    r = _rows_for(dict(rec, city="Hilton Head"), [cand("@address", "3001 Meeting St, Hilton Head Island, SC")])["city"]
    check("city: Hilton Head = Hilton Head Island", r.status == M.CORRECT, r.status)
    r = _rows_for(dict(rec, city="Aiken"), cands)["city"]
    check("city: not among the site's addresses -> Different", r.status == M.DISCREPANCY, r.status)


def test_more_false_alarms():
    cfg = tiny_config()
    page = "Music at St. A. Traditional Anglican choral music is our heritage; a professional organist plays each week."
    copied = {"music_style": ans("Traditional Anglican choral music is our heritage", "Traditional Anglican choral music, with a professional organist"),
              "accessibility": ans("", "unclear", "low")}
    cfg.fields.append(__import__("parishcheck.config", fromlist=["FieldSpec"]).FieldSpec("accessibility", "text", group="community"))
    out = ask(FakeModel(copied), ["music_style", "accessibility"], page, cfg=cfg)
    check("the style example is rejected even when its words are on the page", not out.candidates, str([c.value for c in out.candidates]))
    mixed = {"music_style": ans("Traditional Anglican choral music is our heritage",
                                "Traditional Anglican choral music, with a professional organist; congregational singing of hymns and service music; Adult Choir"),
             "accessibility": ans("", "unclear", "low")}
    out = ask(FakeModel(mixed), ["music_style", "accessibility"], page + " The Adult Choir sings at 10.", cfg=cfg)
    check("copied example items are dropped, the parish's own items kept", [c.value for c in out.candidates] == ["Adult Choir"],
          str([c.value for c in out.candidates]))
    model = PromptModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "service_languages": {"lines": [2], "value": "English", "confidence": "high"}})
    cfg2 = tiny_config()
    cfg2.groups = {"services": {"keywords": ["sunday", "eucharist", "service", "spanish"]}}
    cands, _ = _run_groups(cfg2, model, "Sunday service\n10:00 AM Holy Eucharist\nSpanish lessons on Tuesday")
    eng = [c for c in cands if c.field == "service_languages"]
    check("a model answer of just 'English' counts as the default", len(eng) == 1 and eng[0].method == "rule (default)", str(eng))
    check("doubled clergy title is cleaned", clean_value(cfg.get_field("rector_name"), "The Rev. The Very Reverend Scott Fleischer", cfg)
          == "The Very Rev. Scott Fleischer", clean_value(cfg.get_field("rector_name"), "The Rev. The Very Reverend Scott Fleischer", cfg))
    guess = cand("@email", "hewjr781@me.com", "rector", evidence="contact Fr. Harry Walton in Florida, email: hewjr781@me.com")
    rows = _rows_for(dict(BASE, rector_name="The Rev. Ann Lee"), [guess])
    check("a rector email labelled only by nearby words is low confidence", rows["rector_email"].status == M.UNCLEAR
          and rows["rector_email"].confidence == "low", f"{rows['rector_email'].status} {rows['rector_email'].confidence}")
    named = cand("@email", "alee@church.org", "rector", evidence="The Rev. Ann Lee, Rector: alee@church.org")
    rows = _rows_for(dict(BASE, rector_name="The Rev. Ann Lee"), [named])
    check("... but next to the rector's name it stays high", rows["rector_email"].status == M.SITE_ONLY and rows["rector_email"].confidence == "high",
          f"{rows['rector_email'].status} {rows['rector_email'].confidence}")
    initial = cand("@email", "sfleischer@church.org", "rector", evidence="Email link: sfleischer@church.org")
    rows = _rows_for(dict(BASE), [initial, cand("rector_name", "The Very Rev. Scott Fleischer", method="llm")])
    check("... or when the address holds the rector's surname", rows["rector_email"].confidence == "high", rows["rector_email"].confidence)


def test_rector_house_style():
    real = load_config(HERE / "config.yaml")
    fmt = real.get_field("rector_name").format
    check("rector format hint asks for no role", "no role" in fmt and ", Rector" not in fmt, fmt)
    gold = real.style_examples["rector_name"][0]
    check("style example rector has no role", "," not in gold and "Rector" not in gold, gold)
    sv = merge_field(real.get_field("rector_name"), [cand("@clergy", "The Rev. Marie-Carmel Chery, Dean of the Chapel & Spiritual Engagement")], real)
    check("role dropped from the website value", sv.best == "The Rev. Marie-Carmel Chery", sv.best)
    check("honorifics differ -> same person", persons_match("The Rev. Gary Eichelberger", "Fr. J. Gary Eichelberger")[0]
          and persons_match("The Reverend Jane Doe", "Mtr. Jane Doe")[0])
    check("first name vs initial -> same person", persons_match("The Rev. Jane Smith", "The Rev. J. Smith")[0])
    check("two different initials -> different people", not persons_match("The Rev. J. Smith", "The Rev. M. Smith")[0])
    r = _rows_for(dict(BASE, rector_name="The Rev. Jane Doe"), [cand("rector_name", "The Reverend Jane Doe, Interim Rector", method="llm")])["rector_name"]
    check("rector: only the role differs -> Matches", r.status == M.CORRECT and r.site_value == "The Reverend Jane Doe", f"{r.status} {r.site_value}")


# --------------------------------------------------------------------------------------------------
# Reading the right pages, spotting bad sites (P3)
# --------------------------------------------------------------------------------------------------
class FakeWeb:
    """Stands in for the internet: {url: html}. Any other address is a 404."""

    def __init__(self, pages: dict, redirects: dict | None = None):
        self.pages, self.redirects, self.asked = pages, redirects or {}, []

    def get(self, url, use_cache=True):
        from parishcheck.fetcher import FetchResult
        n = normalize_url(url)
        self.asked.append(n)
        final = self.redirects.get(n, n)
        if final not in self.pages:
            return FetchResult(url=n, final_url=final, status=404, error="HTTP 404")
        return FetchResult(url=n, final_url=final, status=200, content_type="text/html", text=self.pages[final])

    def get_rendered(self, url, **kw):
        from parishcheck.fetcher import FetchResult
        return FetchResult(url=url, error="no browser in the self-test")

    def robots_for(self, url):
        return RobotsRules("")

    def get_sitemap_bytes(self, url):
        return None


def _page(title, body, links=()):
    a = "".join(f'<a href="{h}">{t}</a> ' for h, t in links)
    return f"<html><head><title>{title}</title></head><body><main><h1>{title}</h1><p>{body}</p><nav>{a}</nav></main></body></html>"


FILLER = " Plenty of words about parish life so the page has enough text to read. " * 6


def test_shared_host_crawl():
    from parishcheck.crawler import crawl_site
    from parishcheck.urls import site_key, site_scope
    cfg = tiny_config()
    cfg.use_browser_fallback, cfg.delay_seconds = False, 0
    cfg.shared_host_domains = ["diocese.org", "college.edu"]
    check("parish page on a diocese site gets its own key", site_key("https://diocese.org/st-davids", cfg.shared_host_domains) == "site:diocese.org/st-davids"
          and site_key("https://diocese.org/grace.html", cfg.shared_host_domains) == "site:diocese.org/grace")
    check("ordinary church site: read the whole site", site_scope("https://stjohns.org/", cfg.shared_host_domains) == ("stjohns.org", "")
          and site_scope("https://stjohns.org/home", cfg.shared_host_domains) == ("stjohns.org", ""))
    start = "https://diocese.org/st-davids"
    links = [("/st-davids/worship", "Worship"), ("/staff", "Diocesan staff"), ("/fiscal-affairs", "Finance"), ("/about", "About"),
             ("/news", "News"), ("/st-davids-cheraw-history", "St. David's history"), ("/contact", "Contact")]
    pages = {start: _page("St. David's Episcopal Church", "Sunday worship 10 AM." + FILLER, links),
             "https://diocese.org/st-davids/worship": _page("Worship", "Holy Eucharist 10 AM." + FILLER,
                                                            [("/st-davids/worship/music", "Music"), ("/bishop", "Bishop")]),
             "https://diocese.org/st-davids/worship/music": _page("Music", "Choir." + FILLER)}
    for path in ("/staff", "/fiscal-affairs", "/about", "/news", "/st-davids-cheraw-history", "/contact", "/bishop"):
        pages["https://diocese.org" + path] = _page(path, "Diocesan page." + FILLER, [("/more" + path, "more")])
    web = FakeWeb(pages)
    out = crawl_site(start, cfg, web, __import__("threading").Event(), hint_words=["cheraw", "david"])
    read = sorted(urlsplit_path(d.final_url) for d in out.pages)
    inside = [p for p in read if p.startswith("/st-davids/") or p == "/st-davids"]
    outside = [p for p in read if p not in inside]
    check("shared site: the parish page and the pages below it are read", inside == ["/st-davids", "/st-davids/worship", "/st-davids/worship/music"], str(read))
    check("shared site: at most 3 other pages, only ones the parish page links to", len(outside) <= 3 and "/bishop" not in outside
          and not any(p.startswith("/more") for p in outside), str(outside))
    check("shared site: the page about this parish is among them", "/st-davids-cheraw-history" in outside, str(outside))
    college = {"https://college.edu/": _page("Voorhees University", "St. Philip's Chapel worship." + FILLER,
                                             [(f"/dept{i}", f"Department {i}") for i in range(8)] + [("/chapel", "Chapel")])}
    for i in range(8):
        college[f"https://college.edu/dept{i}"] = _page(f"Dept {i}", "Department." + FILLER, [("/deeper", "x")])
    college["https://college.edu/chapel"] = _page("Chapel", "Sunday worship." + FILLER)
    out = crawl_site("https://college.edu/", cfg, FakeWeb(college), __import__("threading").Event(), hint_words=["philip", "chapel"])
    check("listed shared host: home page plus 3 linked pages", len(out.pages) == 4 and any(d.final_url.endswith("/chapel") for d in out.pages),
          str([d.final_url for d in out.pages]))
    gone = FakeWeb(dict(pages, **{"https://diocese.org/": _page("The Diocese", "Diocesan news." + FILLER)}),
                   redirects={"https://diocese.org/old-parish-page": "https://diocese.org/"})
    out = crawl_site("https://diocese.org/old-parish-page", cfg, gone, __import__("threading").Event(), hint_words=["david"])
    check("shared site: a parish page that now leads to the diocese home is not read as the parish's",
          out.status == "failed" and "is gone" in out.reason and not out.pages, f"{out.status} {out.reason}")


def urlsplit_path(u):
    from urllib.parse import urlsplit
    return urlsplit(u).path or "/"


def test_expired_or_taken_over_sites():
    from parishcheck.crawler import SUSPECT_REASON, crawl_site, suspect_site_reason
    page = lambda html: clean_html(html, "https://x.org/")
    check("parked domain page", suspect_site_reason(page(_page("x.org", "This domain is for sale! Buy this domain today." + FILLER))) == SUSPECT_REASON)
    check("casino spam page", suspect_site_reason(page(_page("Best online casino", "Play poker, roulette and the jackpot at our casino. " * 10))) == SUSPECT_REASON)
    check("spam injected into an old church page", suspect_site_reason(page(_page("St. Andrew's", "Slot gacor and togel. Sportsbook odds." + FILLER)), ["andrew"]) == SUSPECT_REASON)
    check("no church words and no name words", suspect_site_reason(page(_page("Digital marketing", "We grow your brand with SEO and ads. " * 12)), ["andrew"]) == SUSPECT_REASON)
    check("a real church page is fine", suspect_site_reason(page(_page("St. Andrew's", "Join us for worship on Sunday." + FILLER)), ["andrew"]) == "")
    check("a church page mentioning a casino-night fundraiser is fine",
          suspect_site_reason(page(_page("St. Andrew's Episcopal Church", "Casino night fundraiser in the parish hall: poker, blackjack, roulette and a jackpot raffle." + FILLER)), ["andrew"]) == "")
    check("too little text to judge is not flagged", suspect_site_reason(page("<html><body><div id='app'></div></body></html>"), ["andrew"]) == "")
    cfg = tiny_config()
    cfg.use_browser_fallback, cfg.delay_seconds = False, 0
    web = FakeWeb({"https://oldchurch.org/": _page("Lucky Spin Casino", "Online casino, slot gacor and togel. Best betting odds." + FILLER)})
    crawl = crawl_site("https://oldchurch.org/", cfg, web, __import__("threading").Event(), hint_words=["andrew"])
    check("crawl stops at a taken-over home page", crawl.status == "failed" and crawl.site_problem == "suspect" and len(web.asked) == 1,
          f"{crawl.status} {crawl.reason} {web.asked}")
    from parishcheck.extract import ExtractContext, extract_own_site
    model = FakeModel({})
    ctx = ExtractContext(cfg=cfg, client=model, state=_NoState(), today=date(2026, 10, 8), stop=__import__("threading").Event())
    res = extract_own_site("site:oldchurch.org", "https://oldchurch.org/", crawl, ctx)
    check("no model calls for a taken-over site", model.calls == 0 and res["site_problem"] == "suspect")
    rec = dict(BASE, website="https://oldchurch.org/", church_phone="864-235-5884")
    view = build_view(rec, cfg, [res])
    rows = compare_record(rec, cfg, view, "a1", "St. A")
    check("taken-over site: the whole record is Couldn't tell, saying why",
          all(r.status == M.UNCLEAR and r.reason == SUSPECT_REASON for r in rows if r.field not in ("other_contact", "diocesan_bishop")),
          str({r.field: r.status for r in rows}))
    from parishcheck.report import summarize_record
    check("Summary 'Site problem' says so", summarize_record("a1", "St. A", view, rows)["site_problem"] == SUSPECT_REASON)
    # A redirect to another domain: the website field is Different, showing where it goes now.
    web = FakeWeb({"https://newchurch.org/": _page("St. Andrew's Episcopal Church", "Worship Sunday 10 AM." + FILLER)},
                  redirects={"https://oldchurch.org/": "https://newchurch.org/"})
    crawl = crawl_site("https://oldchurch.org/", cfg, web, __import__("threading").Event(), hint_words=["andrew"])
    res = extract_own_site("site:oldchurch.org", "https://oldchurch.org/", crawl, ExtractContext(cfg=cfg, client=None, state=_NoState(),
                           today=date(2026, 10, 8), stop=__import__("threading").Event()))
    view = build_view(rec, cfg, [res])
    r = {x.field: x for x in compare_record(rec, cfg, view, "a1", "St. A")}["website"]
    check("redirect to another domain -> website Different, showing the final address",
          r.status == M.DISCREPANCY and r.site_value == "https://newchurch.org/" and "redirects" in r.evidence, f"{r.status} {r.site_value} {r.evidence}")
    check("... and Summary 'Site problem' names it", "newchurch.org" in view.site_problem, view.site_problem)
    web = FakeWeb({"https://www.church.org/": _page("St. Andrew's", "Worship." + FILLER)}, redirects={"https://church.org/": "https://www.church.org/"})
    crawl = crawl_site("https://church.org/", cfg, web, __import__("threading").Event(), hint_words=["andrew"])
    check("www. redirect is not a different website", crawl.redirected_to == "" and crawl.status == "ok")


def test_out_of_date_pages():
    from parishcheck.llm import stale_hint
    today = date(2026, 10, 9)
    check("a past year next to a date word", stale_hint("Easter 2024: services at 8 and 10", today) == "page may be out of date (mentions 2024)")
    check("summer schedule in October", "summer schedule" in stale_hint("Summer schedule: one service at 9 AM", today))
    check("'through August' in October", "through august" in stale_hint("Through August, one service at 9:00", today).lower())
    check("'beginning June 5' in October", "Beginning June 5" in stale_hint("Beginning June 5 we worship at 9", today))
    check("copyright and history years do not count", stale_hint("© 2019 St. Andrew's. Founded in 1856, rebuilt in 2001.", today) == "")
    check("a school-year schedule does not count", stale_hint("Sunday School, September through May", today) == "")
    check("this year does not count", stale_hint("Christmas 2026 services", today) == "")
    cfg = tiny_config()
    page = "Summer schedule. Sunday worship: 9:00 AM Holy Eucharist. Our choir sings hymns."
    sv = {"sunday_services": ans("Sunday worship: 9:00 AM Holy Eucharist", "9:00 AM Holy Eucharist"),
          "weekday_services": ans("", "unclear", "low"), "service_languages": ans("", "unclear", "low")}
    out = ask(FakeModel(sv), ["sunday_services", "weekday_services", "service_languages"], page, cfg=cfg)
    c = [c for c in out.candidates if c.field == "sunday_services"]
    check("an answer from an out-of-date page is low confidence, with a note", c and c[0].confidence == "low"
          and c[0].extra.get("notes") == ["page may be out of date (mentions a summer schedule)"], str([(x.confidence, x.extra) for x in c]))
    r = _rows_for(dict(BASE, sunday_services="8:00 AM Holy Eucharist"), c)["sunday_services"]
    check("... so a difference is Couldn't tell and the note is shown", r.status == M.UNCLEAR and "out of date" in r.detail, f"{r.status} {r.detail}")
    from parishcheck.extract import _age_bucket
    import time as _t
    fresh, old, unknown = PageDoc("a"), PageDoc("b"), PageDoc("c")
    fresh.modified, old.modified = _t.time() - 30 * 86400, _t.time() - 900 * 86400
    check("recently changed pages are preferred", _age_bucket(fresh, date.today()) == 0 == _age_bucket(unknown, date.today())
          and _age_bucket(old, date.today()) == 2)
    from parishcheck.crawler import _parse_sitemap
    lm = {}
    _parse_sitemap(b"<urlset xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'><url><loc>https://x.org/worship</loc><lastmod>2026-09-01</lastmod></url></urlset>", lm)
    check("sitemap lastmod is read", list(lm) == ["https://x.org/worship"] and lm["https://x.org/worship"] > 0, str(lm))


def test_livestream_prefers_stable_links():
    cfg = tiny_config()
    cfg.fields.append(__import__("parishcheck.config", fromlist=["FieldSpec"]).FieldSpec("livestream_url", "url", group="contact", url_kind="livestream"))
    spec = cfg.get_field("livestream_url")
    video = cand("@livestream", "https://www.youtube.com/watch?v=abc123XYZ", page="https://church.org/a")
    video2 = cand("@livestream", "https://youtu.be/abc123XYZ", page="https://church.org/b")
    channel = cand("@livestream", "https://www.youtube.com/@stjohns/live", page="https://church.org/c")
    sv = merge_field(spec, [video, video2, channel], cfg)
    check("channel /live beats a single video seen on more pages", sv.best == "https://www.youtube.com/@stjohns/live" and not sv.notes, str(sv.values))
    fb = cand("@livestream", "https://www.facebook.com/stjohns/videos", page="https://church.org/d")
    vid = cand("@livestream", "https://www.youtube.com/live/abcdefgh123", page="https://church.org/e")
    check("Facebook /videos page beats youtube.com/live/<id>", merge_field(spec, [vid, fb], cfg).best == "https://www.facebook.com/stjohns/videos")
    sv = merge_field(spec, [video], cfg)
    check("only a single video -> reported with a note", sv.best.startswith("https://www.youtube.com/watch") and
          sv.notes == ["single video link — may go stale; look for the channel URL"], str(sv.notes))
    from parishcheck.normalizers import is_single_video
    check("single-video links recognised", all(is_single_video(u) for u in ["https://youtu.be/abc", "https://www.youtube.com/watch?v=x1",
          "https://www.youtube.com/live/abcdefgh"]) and not any(is_single_video(u) for u in
          ["https://www.youtube.com/@x", "https://www.youtube.com/@x/live", "https://www.youtube.com/channel/UC123", "https://www.facebook.com/x/videos"]))


# --------------------------------------------------------------------------------------------------
# Speed (P4)
# --------------------------------------------------------------------------------------------------
def _run_record_groups(cfg, model, text, record, groups=("services",)):
    from parishcheck.extract import ExtractContext, new_result, run_llm_groups
    ctx = ExtractContext(cfg=cfg, client=model, state=_NoState(), today=date(2026, 10, 8), stop=__import__("threading").Event())
    doc = PageDoc("https://church.org/worship", title="Worship", text=text, main_text=text)
    res = new_result(M.OWN_SITE, "k", "https://church.org/")
    return run_llm_groups([doc], M.OWN_SITE, ctx, list(groups), res, record=record), res


def test_model_skipped_when_data_is_on_the_page():
    cfg = tiny_config()
    cfg.groups = {"services": {"keywords": ["sunday", "eucharist", "service", "wednesday", "a.m.", "noon"]}}
    page = "Sunday services\n8:00 AM Holy Eucharist, Rite I\n10:30 AM Choral Eucharist\nWednesday 12:00 noon Holy Eucharist"
    rec = dict(BASE, sunday_services="8:00 AM Holy Eucharist, Rite I; 10:30 AM Choral Eucharist",
               weekday_services="Wed 12:00 noon Holy Eucharist", service_languages="English")
    model = FakeModel({})
    cands, res = _run_record_groups(cfg, model, page, rec)
    got = {c.field: (c.value, c.method) for c in cands}
    check("JSON filled and its times are on the page -> no model call", model.calls == 0 and res["llm_skipped"] == 1, f"{model.calls} {res}")
    check("... the JSON values come back as rule finds", got.get("sunday_services") == (rec["sunday_services"], "rule")
          and got.get("weekday_services") == (rec["weekday_services"], "rule"), str(got))
    rows = _rows_for(rec, cands)
    check("... and are reported as Matches, method rule", rows["sunday_services"].status == M.CORRECT and rows["sunday_services"].method == "rule"
          and rows["weekday_services"].status == M.CORRECT, f"{rows['sunday_services'].status} {rows['weekday_services'].status}")
    check("... with the line that shows it as evidence", "8:00 AM Holy Eucharist" in rows["sunday_services"].evidence, rows["sunday_services"].evidence)
    blank = PromptModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                         "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"}})
    _run_record_groups(cfg, blank, page, dict(rec, weekday_services=""))
    check("a blank JSON field in the group -> the model is asked", blank.calls == 1)
    moved = FakeModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                       "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"}})
    _run_record_groups(cfg, moved, page, dict(rec, sunday_services="9:00 AM Holy Eucharist"))
    check("a JSON time that is not on the page -> the model is asked", moved.calls == 1)
    shared = FakeModel({"sunday_services": {"lines": [], "value": "unclear", "confidence": "low"},
                        "weekday_services": {"lines": [], "value": "unclear", "confidence": "low"}})
    _run_record_groups(cfg, shared, page, None)
    check("a website shared by several parishes -> the model is asked", shared.calls == 1)
    from parishcheck.pipeline import source_fingerprint
    check("saved results are redone when the JSON values change",
          source_fingerprint("x", cfg, [rec]) != source_fingerprint("x", cfg, [dict(rec, sunday_services="9:00 AM Mass")]))


def test_fields_and_ids_options():
    from parishcheck.config import ConfigError, limit_fields
    from parishcheck.pipeline import read_ids_file, select_indices
    from scraper import build_parser
    args = build_parser().parse_args(["--input", "x.json", "--fields", "rector_name,sunday_services", "--ids-file", "ids.txt"])
    check("--fields and --ids-file are accepted", args.fields == "rector_name,sunday_services" and args.ids_file == "ids.txt")
    cfg = tiny_config()
    limit_fields(cfg, "rector_name, sunday_services")
    rows = _rows_for(dict(BASE, rector_name="The Rev. Ann Lee"), [cand("rector_name", "The Rev. Ann Lee", method="llm")], cfg=cfg)
    check("--fields: only those fields are reported", sorted(rows) == ["rector_name", "sunday_services"], str(sorted(rows)))
    from parishcheck.extract import llm_specs_by_group
    asked = sorted(f.key for g in llm_specs_by_group(cfg).values() for f in g)
    check("--fields: only those fields are asked of the model", asked == ["rector_name", "sunday_services"], str(asked))
    cfg2 = tiny_config()
    limit_fields(cfg2, "rite")
    asked = sorted(f.key for g in llm_specs_by_group(cfg2).values() for f in g)
    check("--fields rite: still reads the services it is built from", asked == ["sunday_services", "weekday_services"], str(asked))
    try:
        limit_fields(tiny_config(), "rector_nmae")
        check("--fields with a typo gives a friendly error", False)
    except ConfigError as exc:
        check("--fields with a typo gives a friendly error", "rector_nmae" in str(exc) and "rector_name" in str(exc), str(exc))
    d = Path(tempfile.mkdtemp(prefix="parish_selftest_ids_"))
    (d / "ids.txt").write_text("# re-check these\ngrace-1\n\n  ch-3  \n", encoding="utf-8")
    ids = read_ids_file(str(d / "ids.txt"))
    check("ids file: one id per line, comments and blanks ignored", ids == ["grace-1", "ch-3"], str(ids))
    recs = [{"id": "st-a", "name": "St. A"}, {"id": "grace-1", "name": "Grace"}, {"id": "grace-2", "name": "Grace Two"}, {"id": "ch-3", "name": "Christ"}]
    check("--ids-file selects exactly those parishes", select_indices(recs, tiny_config(), None, set(ids)) == [1, 3])
    check("--only and --ids-file together", select_indices(recs, tiny_config(), "grace", set(ids)) == [1])


def test_timings_in_run_log():
    import logging as _logging
    from parishcheck.pipeline import _log_timings
    seen = []

    class Grab(_logging.Handler):
        def emit(self, record):
            seen.append(record.getMessage())
    h = Grab()
    lg = _logging.getLogger("parishcheck")
    lg.addHandler(h)
    lg.setLevel(_logging.INFO)
    try:
        saved = {f"site:s{i}.org": {"crawl_seconds": i, "elapsed": i, "llm_calls": 1, "llm_skipped": 1, "pages_crawled": 2} for i in range(15)}
        _log_timings(saved)
    finally:
        lg.removeHandler(h)
    listed = [m for m in seen if m.strip().endswith("model call(s))")]
    check("run.log lists the 10 slowest sites, slowest first", len(listed) == 10 and "site:s14.org" in listed[0], str(listed[:2]))
    check("run.log counts model calls not needed", any("not needed" in m and ": 15;" in m for m in seen), str(seen[:1]))
    check("INSTRUCTIONS.md explains the Excel lock file", "~$Parish Check Results.xlsx" in (HERE / "INSTRUCTIONS.md").read_text(encoding="utf-8"))


# --------------------------------------------------------------------------------------------------
def main() -> int:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    print(f"Running {len(tests)} groups of self-tests (no internet or AI needed)...\n")
    for t in tests:
        before = RESULTS["failed"]
        try:
            t()
            mark = "ok" if RESULTS["failed"] == before else "PROBLEM"
        except Exception:
            RESULTS["failed"] += 1
            mark = "CRASHED"
            print(f"  CRASH in {t.__name__}:")
            traceback.print_exc()
        print(f"  [{mark:^7}] {t.__name__.replace('test_', '').replace('_', ' ')}")
    print(f"\n{RESULTS['passed']} checks passed, {RESULTS['failed']} failed.")
    if RESULTS["failed"]:
        print("SOME TESTS FAILED. Please copy the lines above that start with FAIL and send them to whoever set this up.")
        return 1
    print("ALL TESTS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
