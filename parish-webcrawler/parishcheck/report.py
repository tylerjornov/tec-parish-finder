"""Comparing each record to what was found, and writing the results.

compare_record()        -> one row per field with a status (CORRECT, DISCREPANCY, ...)
write_excel()           -> output/Parish Check Results.xlsx  (for people)
write_results_files()   -> output/results.json, results.csv, suggested_patch.json  (for scripts)
"""

from __future__ import annotations

import csv
import json
import logging
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import models as M
from .comparators import StatusResult, compare_typed, decide_status
from .config import Config, FieldSpec
from .merge import RecordView
from .normalizers import is_missing, norm_text

log = logging.getLogger("parishcheck")

REPORT_COLUMNS = ["id", "name", "city", "field", "status", "json_value", "site_value", "detail", "reason", "method",
                  "source_type", "source_url", "confidence", "evidence"]
# Keys of each object in results.json / columns of results.csv.
RESULT_KEYS = ["id", "name", "city", "field", "status", "label", "json_value", "site_value", "notes", "source_url",
               "method", "confidence", "evidence"]


@dataclass
class Row:
    id: str
    name: str
    field: str
    status: str
    json_value: str = ""
    site_value: str = ""
    detail: str = ""
    reason: str = ""
    method: str = ""
    source_type: str = ""
    source_url: str = ""
    confidence: str = ""
    evidence: str = ""
    city: str = ""
    verified: bool = False         # rule-found, or a model answer whose cited lines hold every fact in it

    def as_list(self) -> list:
        d = asdict(self)
        return [d[c] for c in REPORT_COLUMNS]


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    return str(v)


def _ftype(spec: FieldSpec, cfg: Config) -> str:
    if spec.address_part in ("city", "state"):
        return spec.address_part
    if spec.key == cfg.name_key and spec.type == "text":
        return "name"
    d = (cfg.derived_fields or {}).get(spec.key) or {}
    if "from" in d:
        return "contactlist"
    return spec.type


def record_city(cfg: Config, record: dict) -> str:
    """The city the JSON gives for this record (it tells apart parishes with the same name)."""
    key = next((f.key for f in cfg.fields if f.address_part == "city"), "city")
    v = record.get(key)
    return "" if cfg.is_missing(v) else _fmt(v)


def compare_record(record: dict, cfg: Config, view: RecordView, rid: str, name: str) -> list[Row]:
    """One row per checked field. A problem inside one field becomes an UNCLEAR row; it never stops the run.
    Fields that are never checked (uncheckable_fields, or left out with --fields) get no row."""
    rows: list[Row] = []
    city = record_city(cfg, record)
    for spec in cfg.fields:
        if not spec.checkable or not cfg.reported(spec):
            continue
        try:
            row = (_compare_lookup(record, cfg, rid, name, spec) if spec.lookup
                   else _compare_field(record, cfg, view, rid, name, spec))
        except Exception as exc:                # pragma: no cover - safety net
            log.exception("Could not compare %s for %s", spec.key, rid)
            row = Row(rid, name, spec.key, M.UNCLEAR, _fmt(record.get(spec.key)), "", "",
                      f"internal problem while comparing this field ({type(exc).__name__}: {str(exc)[:100]}); see run.log")
        if row is not None:
            row.city = city
            rows.append(row)
    return rows


def _built_from_json(record: dict, cfg: Config, spec: FieldSpec) -> list[str]:
    """other_contact-style fields: the value the JSON's own parts make ("(864) 235-5884 | office@church.org")."""
    how = (cfg.derived_fields or {}).get(spec.key) or {}
    parts = [_fmt(record.get(k)) for k in how.get("from", []) if not cfg.is_missing(record.get(k))]
    return [how.get("separator", " | ").join(parts)] if parts else []


def _compare_field(record: dict, cfg: Config, view: RecordView, rid: str, name: str, spec: FieldSpec) -> Row:
    key = spec.key
    json_value = record.get(key)
    sv = view.values.get(key)
    ftype = _ftype(spec, cfg)
    sep = ((cfg.derived_fields or {}).get(key) or {}).get("separator", ";")
    if ftype == "contactlist":
        # Built from church_phone and church_email, so compare it with what the JSON's own parts make.
        built = _built_from_json(record, cfg, spec)
        sv = M.SiteValue(values=built, best=built[0] if built else "", method="rule (built from your data)",
                         evidence="built from " + " and ".join(((cfg.derived_fields or {}).get(key) or {}).get("from", [])),
                         verified=True) if built else None
    unclear = ""
    if view.all_failed_reason and ftype != "contactlist":
        unclear = view.all_failed_reason
    elif (sv is None or not sv.values) and spec.type == "url" and not spec.url_kind:
        unclear = "this kind of web-address field is not collected automatically"
    elif (sv is None or not sv.values) and view.issues.get(key):
        unclear = _issue_reason(view.issues[key])
    site_values = list(sv.values) if sv else []
    st: StatusResult = decide_status(
        ftype, json_value, site_values, missing_markers=cfg.missing_markers,
        site_confidence=(sv.confidence if sv else "high"), unclear_reason=unclear,
        json_only_note="; ".join(n for n in (view.limited_note, view.cap_note) if n), default_day=spec.default_day, url_kind=spec.url_kind,
        time_format=cfg.time_format, list_sep=sep,
        json_confirmed_none=cfg.is_confirmed_none(spec, json_value), site_is_default=bool(sv and sv.is_default),
    )
    status, reason, detail_bits = st.status, st.reason, []
    if st.detail:
        detail_bits.append(st.detail)
    if status in (M.DISCREPANCY, M.PARTIAL) or st.json_only or st.site_only:
        if st.json_only:
            detail_bits.append("json_only_items: " + _fmt(st.json_only))
        if st.site_only:
            detail_bits.append("site_only_items: " + _fmt(st.site_only))

    # Phone/email that exist on the site but are labelled for the other role.
    if sv and sv.alt_values and status in (M.DISCREPANCY, M.JSON_ONLY) and not is_missing(json_value, cfg.missing_markers):
        if compare_typed(ftype, str(json_value), sv.alt_values).status == M.CORRECT:
            status, detail_bits = M.CORRECT, [f"found on the site, but labelled as the {'church' if spec.role == 'rector' else 'rector'}'s, not the {spec.role}'s"]

    if status in (M.DISCREPANCY, M.PARTIAL, M.SITE_ONLY) and sv:
        if sv.conflict:
            status, reason = M.UNCLEAR, "the site itself gives different answers: " + " / ".join(sv.values)
        elif view.unlocated or sv.unlocated:
            status, reason = M.UNCLEAR, ("shared website with several locations, and no page showing this record's "
                                         "street address gave this value, so it may belong to another location")
    if sv:
        detail_bits.extend(sv.notes)
    note = "; ".join(n for n in (view.limited_note, view.cap_note) if n)
    if status == M.JSON_ONLY and note and note not in detail_bits:
        detail_bits.append(note)
    elif status in (M.DISCREPANCY, M.PARTIAL, M.SITE_ONLY, M.UNCLEAR) and view.limited_note and view.limited_note not in detail_bits:
        detail_bits.append(view.limited_note)

    site_display = sv.best if sv else ""
    return Row(
        rid, name, key, status, _fmt(json_value), site_display,
        "; ".join(dict.fromkeys(b for b in detail_bits if b)), reason if status == M.UNCLEAR else "",
        (sv.method if sv else ""),
        _source_label(sv.source_type if sv and sv.source_type else view.primary_type),
        (sv.source_url if sv and sv.source_url else view.primary_url),
        (sv.confidence if sv else ""), (sv.evidence if sv else ""),
        verified=bool(sv and sv.verified),
    )


def _lookup_key(text: str) -> str:
    """'The Episcopal Diocese of Upper South Carolina' -> 'upper south carolina'."""
    return " ".join(w for w in norm_text(text).split() if w not in {"the", "episcopal", "diocese", "of", "church", "in"})


def _compare_lookup(record: dict, cfg: Config, rid: str, name: str, spec: FieldSpec) -> Optional[Row]:
    """diocesan_bishop: checked against diocese_lookup in config.yaml (diocese -> current bishop), not crawled."""
    table = cfg.diocese_lookup
    if not table:
        return None
    json_value = record.get(spec.key)
    diocese = record.get(spec.lookup)
    if cfg.is_missing(diocese):
        return Row(rid, name, spec.key, M.UNCLEAR, _fmt(json_value), "",
                   reason=f"the record has no {spec.lookup}, so its bishop cannot be looked up")
    bishop = next((b for d, b in table.items() if _lookup_key(d) == _lookup_key(str(diocese))), None)
    if bishop is None:
        return Row(rid, name, spec.key, M.UNCLEAR, _fmt(json_value), "",
                   reason=f"the {spec.lookup} '{diocese}' is not in diocese_lookup in config.yaml")
    st = decide_status(spec.type, json_value, [bishop], missing_markers=cfg.missing_markers,
                       json_confirmed_none=cfg.is_confirmed_none(spec, json_value))
    detail = "; ".join(x for x in (st.detail, f"current bishop of {diocese} according to diocese_lookup in config.yaml") if x)
    return Row(rid, name, spec.key, st.status, _fmt(json_value), bishop, detail, st.reason,
               "rule (diocese lookup)", "config.yaml", "", "high", f"diocese_lookup: {diocese} -> {bishop}", verified=True)


def _source_label(kind: str) -> str:
    return {"own_site": M.OWN_SITE, M.OWN_SITE: M.OWN_SITE, "facebook": M.FACEBOOK, M.FACEBOOK: M.FACEBOOK,
            "asset_map": M.ASSET_MAP, M.ASSET_MAP: M.ASSET_MAP}.get(kind, kind or "")


def _issue_reason(issues: list[str]) -> str:
    if any(i.startswith("rejected_unverified") for i in issues):
        return ("the model gave an answer but it could not be verified in the page text, so I threw it away "
                f"({issues[0].split(': ', 1)[-1][:140]})")
    return "the model could not give a readable answer for this field"


def summarize_record(rid: str, name: str, view: RecordView, rows: list[Row]) -> dict:
    counts = Counter(r.status for r in rows)
    notes = list(dict.fromkeys(n for n in view.notes if n))
    for n in (view.limited_note, view.cap_note):
        if n:
            notes.append(n)
    if view.all_failed_reason:
        notes.insert(0, view.all_failed_reason)
    return {
        "id": rid, "name": name, "sources_used": " | ".join(view.sources_used) or "(none)",
        "pages_crawled": view.pages_crawled, "cap_hit": view.cap_hit, "counts": dict(counts),
        "notes": " ; ".join(notes)[:900], "site_problem": view.site_problem,
    }


# --------------------------------------------------------------------------------------
# The results workbook
# --------------------------------------------------------------------------------------
RESULTS_NAME = "Parish Check Results"

# status -> (plain-English label, fill colour, what it means)
STATUS_STYLE = {
    M.DISCREPANCY:   ("Different on website", "F8D7D5", "Your data and the website both have a value, and they differ."),
    M.PARTIAL:       ("Partly matches", "FCE8C3", "Some items match and some don't. The notes list which."),
    M.SITE_ONLY:     ("Missing from your data", "D6E6FA", "The website has it, but your data is blank. 'Website says' is the value to add."),
    M.JSON_ONLY:     ("Not found on website", "E8DEF5", "Your data has it, but the website doesn't mention it. It may be out of date."),
    M.UNCLEAR:       ("Couldn't tell", "E7E7E7", "The tool could not decide. The notes say why."),
    M.CORRECT:       ("Matches", "DCEFD9", "Your data matches the website."),
    M.BOTH_MISSING:  ("Missing in both", "FFFFFF", "Neither your data nor the website has it."),
    M.NOT_CHECKABLE: ("Not checked", "FFFFFF", "A field this tool skips (id, lat, lon, notes...)."),
}

_HEADER_FILL = PatternFill("solid", fgColor="2F3E4E")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(wrap_text=True, vertical="top")
_GROUP_TOP = Border(top=Side(style="medium", color="7F8C99"))
_GREY = Font(color="9AA3AD")
_BOLD = Font(bold=True)
_LINK = Font(color="1F5FBF", underline="single")


def _label(field_key: str) -> str:
    return field_key.replace("_", " ").strip().capitalize()


def status_label(status: str) -> str:
    return STATUS_STYLE.get(status, (status,))[0]


def _clean(v) -> str:
    v = ILLEGAL_CHARACTERS_RE.sub("", _fmt(v))
    return v[:32000]


def _set(ws, row: int, col: int, value, font=None, fill=None):
    c = ws.cell(row=row, column=col)
    text = _clean(value)
    c.value = text
    if text.startswith("="):
        c.data_type = "s"           # never let Excel treat site text as a formula
    c.alignment = _WRAP
    if font is not None:
        c.font = font
    if fill is not None:
        c.fill = fill
    return c


def _link(ws, row: int, col: int, url: str, fill=None):
    c = _set(ws, row, col, url, fill=fill)
    if url.startswith(("http://", "https://")):
        c.hyperlink = url
        c.font = _LINK
    return c


def _sheet(wb, title: str, headers: list[str], widths: list[int], header_row: int = 1):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"B{header_row + 1}"
    ws.row_dimensions[header_row].height = 22
    return ws


def _notes(r: Row) -> str:
    return "; ".join(x for x in (r.reason, r.detail) if x)


def _check_rows(ws, rows: list[Row], extra) -> None:
    """Rows grouped by parish: the first row of each parish is bold with a line above it."""
    prev = None
    for n, r in enumerate(rows, start=2):
        label, colour, _ = STATUS_STYLE.get(r.status, (r.status, "FFFFFF", ""))
        fill = PatternFill("solid", fgColor=colour)
        first = (r.id != prev)
        prev = r.id
        _set(ws, n, 1, r.name or r.id, font=_BOLD if first else _GREY)
        _set(ws, n, 2, r.id, font=None if first else _GREY)
        _set(ws, n, 3, _label(r.field))
        _set(ws, n, 4, label, fill=fill, font=_BOLD)
        _set(ws, n, 5, r.json_value)
        _set(ws, n, 6, r.site_value)
        _set(ws, n, 7, _notes(r))
        _link(ws, n, 8, r.source_url)
        for i, v in enumerate(extra(r), start=9):
            _set(ws, n, i, v)
        if first and n > 2:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=n, column=col).border = _GROUP_TOP
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{len(rows) + 1}"


def _by_parish(rows: list[Row]) -> list[Row]:
    order = {st: i for i, st in enumerate(STATUS_STYLE)}
    return sorted(rows, key=lambda r: ((r.name or "").casefold(), str(r.id), order.get(r.status, 99)))


def meta_line(meta: Optional[dict]) -> str:
    if not meta:
        return ""
    bits = [f"Run {meta.get('run_at', '')}", f"file {meta.get('input_file', '')}", f"model {meta.get('model', '')}"]
    if meta.get("fields"):
        bits.append("fields " + ", ".join(meta["fields"]))
    if meta.get("ids_file"):
        bits.append(f"ids from {meta['ids_file']}")
    if meta.get("only"):
        bits.append(f"only '{meta['only']}'")
    return " · ".join(b for b in bits if b)


def write_excel(cfg: Config, rows: list[Row], summary: list[dict], meta: Optional[dict] = None) -> Path:
    """Write output/Parish Check Results.xlsx and return its path."""
    wb = Workbook()
    wb.remove(wb.active)
    cols = ["Parish", "Id", "Field", "Problem", "Your data", "Website says", "Notes", "Source page"]
    widths = [30, 22, 18, 22, 40, 40, 45, 35]

    review = _by_parish([r for r in rows if r.status in M.REVIEW_STATUSES])
    ws = _sheet(wb, "To review", cols + ["Confidence", "Evidence (words on the page)"], widths + [12, 60])
    _check_rows(ws, review, lambda r: [r.confidence, r.evidence])

    head = 2 if meta else 1
    ws = _sheet(wb, "Summary", ["Parish", "Id", "Website", "Site problem", "To review", "Matches", "Couldn't tell",
                                "Pages read", "Notes"], [34, 22, 35, 30, 11, 10, 13, 11, 70], header_row=head)
    if meta:
        _set(ws, 1, 1, meta_line(meta), font=_BOLD)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    review_fill = PatternFill("solid", fgColor=STATUS_STYLE[M.DISCREPANCY][1])
    ok_fill = PatternFill("solid", fgColor=STATUS_STYLE[M.CORRECT][1])
    problem_fill = PatternFill("solid", fgColor=STATUS_STYLE[M.UNCLEAR][1])
    ordered = sorted(summary, key=lambda s: ((s["name"] or "").casefold(), str(s["id"])))
    for n, s in enumerate(ordered, start=head + 1):
        c = s["counts"]
        to_review = sum(c.get(st, 0) for st in M.REVIEW_STATUSES)
        notes = s["notes"]
        if s.get("pdfs"):
            notes = (notes + " ; " if notes else "") + "PDFs found (not read): " + ", ".join(s["pdfs"][:10])
        _set(ws, n, 1, s["name"] or s["id"], font=_BOLD)
        _set(ws, n, 2, s["id"])
        _link(ws, n, 3, s.get("website", ""))
        _set(ws, n, 4, s.get("site_problem", ""), fill=problem_fill if s.get("site_problem") else None)
        ws.cell(row=n, column=5, value=to_review).fill = review_fill if to_review else ok_fill
        ws.cell(row=n, column=6, value=c.get(M.CORRECT, 0))
        ws.cell(row=n, column=7, value=c.get(M.UNCLEAR, 0))
        ws.cell(row=n, column=8, value=s["pages_crawled"])
        _set(ws, n, 9, notes)
    if summary:
        ws.auto_filter.ref = f"A{head}:I{len(summary) + head}"

    ws = _sheet(wb, "All checks", cols[:3] + ["Result"] + cols[4:] + ["Found by", "Confidence", "Evidence (words on the page)"],
                widths + [10, 12, 60])
    _check_rows(ws, _by_parish(rows), lambda r: [r.method, r.confidence, r.evidence])

    ws = _sheet(wb, "Key", ["Label", "What it means"], [26, 90])
    ws.freeze_panes = "A2"
    shown = [v for k, v in STATUS_STYLE.items() if k != M.NOT_CHECKABLE]
    for n, (label, colour, meaning) in enumerate(shown, start=2):
        _set(ws, n, 1, label, font=_BOLD, fill=PatternFill("solid", fgColor=colour))
        _set(ws, n, 2, meaning)
    n = len(shown) + 3
    for line in ("'To review' lists only the rows that need a look. 'All checks' has every checked field of every parish.",
                 "Your JSON file was not changed. Edit it yourself using 'Website says' and the source page link.",
                 "The AI makes mistakes: treat this as a to-do list, and check the source page before changing anything.",
                 "The same results are in results.json and results.csv (for scripts). suggested_patch.json holds only "
                 "the most certain changes."):
        _set(ws, n, 1, line)
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=2)
        n += 1

    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    path = cfg.output_dir / f"{RESULTS_NAME}.xlsx"
    try:
        wb.save(path)
    except PermissionError:             # the old results are open in Excel
        path = cfg.output_dir / f"{RESULTS_NAME} (new).xlsx"
        wb.save(path)
    return path


# --------------------------------------------------------------------------------------
# Files for scripts: results.json, results.csv, suggested_patch.json
# --------------------------------------------------------------------------------------
def result_dict(r: Row) -> dict:
    return {"id": r.id, "name": r.name, "city": r.city, "field": r.field, "status": r.status,
            "label": status_label(r.status), "json_value": r.json_value, "site_value": r.site_value,
            "notes": _notes(r), "source_url": r.source_url, "method": r.method, "confidence": r.confidence,
            "evidence": r.evidence}


def is_patch_candidate(r: Row) -> bool:
    """High-signal changes only: a value missing from the data or different on the site, found with high
    confidence by a rule or by the model with every fact in its cited lines."""
    return (r.status in (M.SITE_ONLY, M.DISCREPANCY) and r.confidence == "high"
            and (r.method.startswith("rule") or r.verified))


def _in_field_order(cfg: Config, rows: list[Row]) -> list[Row]:
    pos = {f.key: i for i, f in enumerate(cfg.fields)}
    return sorted(rows, key=lambda r: ((r.name or "").casefold(), str(r.id), pos.get(r.field, 999)))


def write_results_files(cfg: Config, rows: list[Row], meta: dict) -> list[Path]:
    """results.json ({"meta": ..., "results": [one object per field check]}), the same rows as results.csv, and
    suggested_patch.json ([{id, field, old, new, source_url, evidence}]). UTF-8, nothing cut short."""
    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    ordered = _in_field_order(cfg, [r for r in rows if r.status != M.NOT_CHECKABLE])
    results = [result_dict(r) for r in ordered]
    meta = dict(meta, counts=dict(Counter(r["status"] for r in results)))

    paths = [cfg.output_dir / "results.json", cfg.output_dir / "results.csv", cfg.output_dir / "suggested_patch.json"]
    paths[0].write_text(json.dumps({"meta": meta, "results": results}, ensure_ascii=False, indent=2), encoding="utf-8")
    with open(paths[1], "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=RESULT_KEYS)
        w.writeheader()
        w.writerows(results)
    patch = [{"id": r.id, "field": r.field, "old": r.json_value, "new": r.site_value, "source_url": r.source_url,
              "evidence": r.evidence} for r in ordered if is_patch_candidate(r)]
    paths[2].write_text(json.dumps(patch, ensure_ascii=False, indent=2), encoding="utf-8")
    return paths
