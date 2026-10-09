#!/usr/bin/env python3
"""parish_check.py - checks data/parishes.json against each parish's own website, with fixed rules (no AI).

Run:        double-click run.command (asks which JSON file)   or   .venv/bin/python parish_check.py path/to/file.json
Self-test:  .venv/bin/python parish_check.py --selftest      (offline, a few seconds)

How it works
  1. Reads each parish's website: up to 30 pages, contact/staff/about pages first, politely (robots.txt, one
     request per second per site, pages cached for 14 days in output/.cache). A parish whose website is an
     Episcopal Asset Map listing gets that listing read, then the church website the listing links to.
  2. Finds facts with fixed rules: phone numbers and emails (labelled church or rector by the words next to
     them), street addresses, livestream links, and the lead priest ("The Rev. Ann Lee" above "Rector").
  3. Compares them with the JSON and keeps only changes it is sure of. A value that is merely missing from the
     website is never reported, and nothing is ever deleted.

Output (output/):
  corrections.json        every change to make to parishes.json, written for a script or an AI to apply
  Corrections.xlsx        the same changes, for people
  Website Problems.xlsx   parishes whose website could not be read, and why (colour-coded)
parishes.json itself is never changed.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import heapq
import json
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit
from urllib.robotparser import RobotFileParser

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ==================================================================================================
# Settings
# ==================================================================================================
HERE = Path(__file__).resolve().parent
OUT = HERE / "output"
CACHE = OUT / ".cache"

MAX_PAGES, MAX_DEPTH = 30, 3          # pages read per website, and how many clicks deep
DELAY, WORKERS, TIMEOUT = 1.0, 6, 20  # seconds between requests to one site, sites read at once, seconds per request
CACHE_DAYS = 14
USER_AGENT = "Mozilla/5.0 (compatible; ParishInfoVerifier/2.0; low-volume parish data check; respects robots.txt)"

BLANK = {"", "n/a", "n/a - data not available"}          # values that count as missing (any capitals)
SHARED_HOSTS = ("episcopalchurchsc.org", "edusc.org", "voorhees.edu")   # one site, many parishes
BISHOPS = {"Upper South Carolina": "The Rt. Rev. Daniel P. Richards",  # keep up to date: diocese -> bishop
           "South Carolina": "The Rt. Rev. Ruth Woodliff-Stanley"}

GOOD_PAGE = re.compile(r"contact|staff|clergy|leader|about|who-we-are|our-team|people|rector|vicar|visit|welcome|"
                       r"location|direction|worship|live|watch|stream", re.I)
STAFF_PAGE = re.compile(r"staff|clergy|leader|who-we-are|our-team|people|rector|vicar|ministers|about|contact", re.I)
NOT_TODAY = re.compile(r"bishop|visitation|history", re.I)    # clergy named on these pages are not today's rector
SKIP_URL = re.compile(r"/(?:19|20)\d\d/|calendar|/events?/|/tags?/|/category/|/author/|/feed|wp-json|wp-admin|login|"
                      r"logout|/cart|/checkout|/shop|/product|/page/\d|[?&](?:s|ical|share|replytocom)=|"
                      r"\.(?:pdf|jpe?g|png|gif|svg|webp|mp[34]|mov|zip|docx?|xlsx?|pptx?|ics|xml|css|js)(?:$|\?)", re.I)

# Website problems: label -> (colour, meaning). Order = order in the Excel file.
PROBLEMS = {
    "Blocked":               ("F8D7D5", "The website refuses automated visitors (HTTP 401, 403 or 429). Check it in a browser."),
    "Not found":             ("FCE8C3", "The address leads nowhere: page not found (404), or the parish's page on a shared site is gone."),
    "Can't connect":         ("FFF2CC", "No answer: the domain doesn't exist, the server is down or slow, or its security certificate is broken."),
    "Expired or taken over": ("E8DEF5", "The address now shows a parked, for-sale, spam or unrelated website."),
    "Not readable":          ("D6E6FA", "Only a Facebook page, a robots.txt that forbids visits, or a page with no readable text."),
    "Shared site":           ("DCEFD9", "A diocese or college website shared by several parishes, with no page of this parish's own to read."),
    "No website":            ("E7E7E7", "parishes.json has no website for this parish."),
}


# ==================================================================================================
# Small text helpers
# ==================================================================================================
def ws(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip().casefold() in BLANK)


def domain(url: str) -> str:
    """'https://www.Church.org/x' -> 'church.org'."""
    host = urlsplit(url if "//" in url else "//" + url).hostname or ""
    return host.lower().removeprefix("www.")


def home_url(url: str) -> str:
    p = urlsplit(url)
    return f"{p.scheme}://{p.netloc}/"


PHONE_RE = re.compile(r"(?<![\d.\-/])(?:\+?1[\s.\-]?)?(?:\(\s*([2-9]\d\d)\s*\)\s*|([2-9]\d\d)[\s.\-]+)([2-9]\d\d)[\s.\-]+(\d{4})(?![\d\-])")


def phone_digits(text: str) -> list[str]:
    out = [(a or b) + c + d for a, b, c, d in PHONE_RE.findall(text or "")]
    if not out:     # tel: links and JSON values may have no separators at all
        d = re.sub(r"\D", "", text or "")
        d = d[1:] if len(d) == 11 and d[0] == "1" else d
        out = [d] if len(d) == 10 and d[0] in "23456789" else []
    return [d for d in out if d[3:6] != "555"]      # 555 numbers are website-template placeholders


def fmt_phone(d: str) -> str:
    return f"({d[:3]}) {d[3:6]}-{d[6:]}"


EMAIL_RE = re.compile(r"(?<![\w.+-])[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,24}(?![\w-])")
JUNK_EMAIL = re.compile(r"\.(?:png|jpe?g|gif|webp|svg|css|js)$|@(?:sentry|wixpress|example|domain|email|yourdomain|"
                        r"godaddy|squarespace|mysite|mailservice|company|website)\.|^(?:mymail|yourname|name|email|user)@", re.I)
# Mailboxes for one ministry or person, never the church's main email.
NOT_MAIN_EMAIL = re.compile(r"financ|giving|give|donat|pledge|treasur|bookkeep|account|music|choir|organ|youth|child|"
                            r"school|preschool|formation|outreach|prayer|webmaster|news|communicat|altar|flower|events", re.I)


def emails_in(text: str) -> list[str]:
    return [e.lower() for e in EMAIL_RE.findall(text or "") if not JUNK_EMAIL.search(e)]


# ==================================================================================================
# Addresses
# ==================================================================================================
STATES = {"alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA", "colorado": "CO",
          "connecticut": "CT", "delaware": "DE", "district of columbia": "DC", "florida": "FL", "georgia": "GA",
          "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
          "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI",
          "minnesota": "MN", "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
          "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY", "north carolina": "NC",
          "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR", "pennsylvania": "PA",
          "rhode island": "RI", "south carolina": "SC", "south dakota": "SD", "tennessee": "TN", "texas": "TX",
          "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA", "west virginia": "WV",
          "wisconsin": "WI", "wyoming": "WY"}
SUFFIX = {"street": "St", "st": "St", "avenue": "Ave", "ave": "Ave", "av": "Ave", "road": "Rd", "rd": "Rd",
          "drive": "Dr", "dr": "Dr", "boulevard": "Blvd", "blvd": "Blvd", "lane": "Ln", "ln": "Ln", "highway": "Hwy",
          "hwy": "Hwy", "court": "Ct", "ct": "Ct", "place": "Pl", "pl": "Pl", "circle": "Cir", "cir": "Cir",
          "parkway": "Pkwy", "pkwy": "Pkwy", "square": "Sq", "sq": "Sq", "terrace": "Ter", "ter": "Ter",
          "trail": "Trl", "trl": "Trl", "way": "Way", "pike": "Pike", "route": "Rte", "rte": "Rte", "plaza": "Plaza",
          "row": "Row", "loop": "Loop", "alley": "Aly"}
DIRS = {"north": "N", "south": "S", "east": "E", "west": "W", "northeast": "NE", "northwest": "NW",
        "southeast": "SE", "southwest": "SW", **{d: d.upper() for d in ("n", "s", "e", "w", "ne", "nw", "se", "sw")}}
_STATE_ALT = "|".join(sorted((s.title() for s in STATES), key=len, reverse=True))
ADDRESS_RE = re.compile(
    r"(?<![\w/#.\-–])(?P<num>\d{1,6}[A-Za-z]?)\s+(?!\d+\s)"
    r"(?P<street>(?:(?:N|S|E|W|NE|NW|SE|SW|North|South|East|West)\.?\s+)?(?:[A-Z0-9][\w'’.\-]*\s+){0,4}?"
    r"(?i:" + "|".join(SUFFIX) + r")\b\.?(?:\s+\d{1,4}[A-Za-z]?\b)?(?:\s+(?:NE|NW|SE|SW|N|S|E|W)\b\.?)?)"
    r"(?:\s*,?\s*(?i:suite|ste|unit|room|bldg)\.?\s*[\w-]+)?\s*(?:[,\n|•·]\s*|\s+)"
    r"(?P<city>[A-Z][a-z][A-Za-z.'’\-]*(?:\s+[A-Z][a-z][A-Za-z.'’\-]*){0,3})[\s,]+(?P<state>[A-Z]{2}|" + _STATE_ALT + r")\b")


def state_code(s: str) -> str:
    s = ws(s).strip(".")
    return s.upper() if s.upper() in STATES.values() else STATES.get(s.lower(), "")


def fmt_street(street: str) -> str:
    """House style: '401 North Main Street' -> '401 N Main St'."""
    out = []
    for t in ws(street).split():
        low = t.lower().strip(".,")
        t = DIRS.get(low) or SUFFIX.get(low) or (t.title() if t.isupper() and len(t) > 2 else t)
        out.append(t)
    return " ".join(out)


def find_addresses(text: str) -> list[tuple[str, str]]:
    """[('401 N Main St, Allendale, SC', matched text)] in house style."""
    text = PHONE_RE.sub(lambda m: "\n" + " " * (len(m.group(0)) - 1), text or "")  # "Call 843-681-8333 3001 Main St"
    out = {}
    for m in ADDRESS_RE.finditer(text):
        st = state_code(m.group("state"))
        if st:
            a = f"{fmt_street(m.group('num') + ' ' + m.group('street'))}, {ws(m.group('city'))}, {st}"
            out.setdefault(a, ws(m.group(0)))
    return list(out.items())


def _street_key(addr: str) -> tuple[str, frozenset]:
    street = addr.split(",")[0]
    toks = [t.lower().strip(".") for t in street.split()]
    if not toks or not re.match(r"\d", toks[0]):
        return "", frozenset()
    core = {t for t in toks[1:] if t not in SUFFIX and t not in DIRS}
    return toks[0], frozenset(core)


def addresses_match(a: str, b: str) -> bool:
    """Same street number and street name; 'Street'/'St', 'North'/'N' and the city are ignored."""
    (na, ca), (nb, cb) = _street_key(a), _street_key(b)
    return bool(na) and na == nb and bool(ca and cb) and (ca <= cb or cb <= ca)


def split_address(addr: str) -> tuple[str, str]:
    parts = [p.strip() for p in addr.split(",")]
    return (parts[1] if len(parts) >= 3 else ""), (state_code(parts[-1]) if len(parts) >= 2 else "")


def city_key(c: str) -> str:
    return " ".join({"mt": "mount", "st": "saint", "ft": "fort"}.get(w, w) for w in re.findall(r"[a-z]+", (c or "").lower()))


# ==================================================================================================
# People
# ==================================================================================================
HONORIFIC = {"the", "reverend", "rev", "revd", "very", "right", "rt", "most", "fr", "father", "mother", "mtr", "dr",
             "canon", "deacon", "bishop", "mr", "mrs", "ms", "pastor"}
SUFFIXES = {"jr", "sr", "ii", "iii", "iv", "phd", "dmin"}
TITLES = [(re.compile(r"\b(?:Right|Rt\.?)\s*(?:Reverend|Revd\.?|Rev\b\.?)", re.I), "Rt. Rev."),
          (re.compile(r"\bVery\s+(?:Reverend|Revd\.?|Rev\b\.?)", re.I), "Very Rev."),
          (re.compile(r"\b(?:Reverend|Revd\.?|Rev\b\.?)", re.I), "Rev."),
          (re.compile(r"\b(?:Father|Fr\b\.?)(?=\s+[A-Z])"), "Fr."),
          (re.compile(r"\b(?:Mother|Mtr\b\.?)(?=\s+[A-Z])"), "Mtr.")]


def person(raw: str) -> tuple[list[str], str]:
    """'The Reverend J. Gary Eichelberger, Rector' -> (['j', 'gary'], 'eichelberger')."""
    head = re.split(r"\s*(?:,|\(|\s[-–—]\s)\s*", ws(raw))[0]
    toks = re.findall(r"[a-z][a-z'\-]*", head.lower().replace("’", "'"))
    while toks and toks[0] in HONORIFIC:
        toks.pop(0)
    while toks and toks[-1] in SUFFIXES:
        toks.pop()
    return (toks[:-1], toks[-1]) if toks else ([], "")


def same_person(a: str, b: str) -> bool:
    (ga, sa), (gb, sb) = person(a), person(b)
    if not sa or not sb or not (sa == sb or set(sa.split("-")) & set(sb.split("-"))):
        return False
    if ga and gb:    # first names or initials must fit together ("J. Gary" = "Gary", "Jane" != "John")
        return any(x[0] == y[0] if min(len(x), len(y)) == 1 else (x.startswith(y) or y.startswith(x)) for x in ga for y in gb)
    return True


def house_title(name: str) -> str:
    """House style: 'REVEREND JANE DOE' -> 'The Rev. Jane Doe'."""
    words = []
    for t in ws(name).split():
        if t.isupper() and len(t) > 2 or t.islower() or re.fullmatch(r"[A-Z]{2}[a-z]+", t):
            t = "-".join(p[:1].upper() + p[1:].lower() for p in t.split("-"))
        words.append(t)
    name = " ".join(words)
    for rx, rep in TITLES:
        name = rx.sub(rep, name)
    name = re.sub(r"^the\s+", "", name, flags=re.I)
    return f"The {name}" if re.match(r"(?:Rt\. Rev\.|Very Rev\.|Rev\.)", name) else name


# A role on a line of its own, and what kind of clergy it is.
ROLE_WORDS = r"(?:rector|vicar|priest|dean|curate|deacon|bishop|canon|chaplain|pastor|seminarian|missioner)"
ROLE_START = re.compile(r"^(?:(?:meet\s+)?our\s+|meet\s+the\s+|the\s+)?(?:(?:interim|associate|assistant|assisting|"
                        r"acting|senior|retired|vocational|transitional|honorary|supply)\s+)*" + ROLE_WORDS + r"\b", re.I)
ROLE_TAIL = re.compile(r"^(?:[\s,&/\-–]*(?:and\s+)?(?:emerit\w*|retired|interim|elect|associate|in[\s-]+charge|"
                       r"in[\s-]+residence|on[\s-]+call|(?:of|for|to)\s.*|" + ROLE_WORDS + r"))*\s*$", re.I)
NOT_ROLE = re.compile(r"\bto the\b|['’]s\b|\d|@|www\.|https?:|:|\b(?:rev|reverend|fr|father|mother|mtr)\b|warden|"
                      r"director|administrator|secretary|coordinator|manager|treasurer|organist|corner|message|search|"
                      r"welcome|blog|letter|note|sermon", re.I)
STAFF_TITLE = re.compile(r"administrator|secretary|director|coordinator|manager|treasurer|organist|warden|clerk|sexton|"
                         r"verger|assistant|minister|bookkeeper|musician|nurse", re.I)
NAME_LINE = re.compile(r"^(?P<hon>(?:the\s+)?(?:(?:very|right|rt\.?|most)\s+)?(?:rev(?:erend|d)?\.?|fr\.?|father|"
                       r"mother|mtr\.?|bishop)(?:\s+(?:dr|canon)\.?)*\s+)?(?P<name>[A-Za-z][\w.'\-]*(?:\s+[A-Za-z][\w.'\-]*){1,4})"
                       r"(?P<sfx>(?:,?\s+(?:jr|sr|ii|iii|iv|ph\.?\s?d|d\.?\s?min)\.?)*)$", re.I)
NOT_NAME = {"us", "our", "staff", "clergy", "contact", "the", "of", "and", "for", "to", "welcome", "meet", "leadership",
            "vestry", "office", "email", "phone", "ministry", "ministries", "music", "church", "parish", "episcopal",
            "director", "team", "about", "worship", "services", "home", "give", "events", "news", "bio", "profile",
            "sermon", "sermons", "rector", "vicar", "priest", "dean", "deacon", "bishop", "curate", "canon"}


def clergy_kind(role: str) -> str:
    """'lead' (Rector, Interim Rector, Vicar, Priest-in-Charge, Dean), 'weak' (just "Priest"), 'bishop' or 'other'."""
    low = ws(role).lower()
    if "bishop" in low:
        return "bishop"
    if re.search(r"emerit|retired|associate|assistant|assisting|curate|deacon|seminarian|residence|chaplain|"
                 r"honorary|supply|former|visiting", low):
        return "other"
    if re.search(r"\b(?:rector|vicar|dean)\b|priest[\s-]*in[\s-]*charge|interim\s+priest", low):
        return "lead"
    return "weak" if re.fullmatch(r"(?:our |the |parish )?priest", low) else "other"


def role_line(line: str) -> str:
    line = ws(line).strip(" .,;-–—")
    m = ROLE_START.match(line)
    return line if m and len(line) <= 60 and not NOT_ROLE.search(line) and ROLE_TAIL.match(line[m.end():]) else ""


def name_line(line: str) -> str:
    """The person's name when the whole line is one, '-' for a vacant post, else ''."""
    line = ws(line).strip(" ,.;:-")
    if re.fullmatch(r"vacant|tba|tbd|open|position open|search in progress", line, re.I):
        return "-"
    line = re.sub(r"\s*(?:[“\"][^”\"]{1,20}[”\"]|\([A-Za-z]{2,15}\))\s*", " ", line).strip()   # “Janey”, (Jim)
    m = NAME_LINE.match(line) if len(line) <= 70 else None
    if not m:
        return ""
    toks = m.group("name").split()
    if any(t.lower().strip(".") in NOT_NAME or re.search(r"['’]s$", t) for t in toks):
        return ""
    if not m.group("hon") and not all(t[:1].isupper() for t in toks):
        return ""
    return ws(f"{m.group('hon') or ''} {m.group('name')}{m.group('sfx') or ''}")


def find_clergy(lines: list[str]) -> list[tuple[str, str, str]]:
    """[(name, role, evidence)]: a role on its own line belongs to the name above it, or below it when that block
    of name/title lines starts with a title; also one-line forms "The Rev. Ann Lee, Rector" / "Rector: ..."."""
    names, roles = [name_line(x) for x in lines], [role_line(x) for x in lines]
    kind = ["N" if names[i] and not roles[i] else "T" if roles[i] or (len(x) <= 60 and STAFF_TITLE.search(x)) else ""
            for i, x in enumerate(lines)]
    out, block = [], ""
    for i, line in enumerate(lines):
        if kind[i] and (i == 0 or not kind[i - 1]):
            block = kind[i]
        if roles[i]:
            j = i + 1 if block == "T" or re.match(r"(?:meet\s+)?(?:our|the)\s", roles[i], re.I) else i - 1
            if 0 <= j < len(lines) and kind[j] == "N" and names[j] != "-":
                out.append((names[j], roles[i], f"{lines[min(i, j)]} / {lines[max(i, j)]}"))
        elif not names[i] and len(line) <= 100:
            m = re.match(r"^(.+?)\s*(?:,|\s[-–—|]\s|:|\()\s*([^()]+?)\)?$", ws(line))
            for n, r in ((m.group(1), m.group(2)), (m.group(2), m.group(1))) if m else ():
                nm, rl = name_line(n), role_line(r)
                if nm and nm != "-" and rl and (re.match(r"(?i)the|rev|fr|father|mother|mtr", nm) or clergy_kind(rl) != "other"):
                    out.append((nm, rl, ws(line)))
                    break
    return out


# ==================================================================================================
# Fetching (polite, cached) and reading pages
# ==================================================================================================
class Web:
    def __init__(self):
        self.client = httpx.Client(headers={"User-Agent": USER_AGENT}, follow_redirects=True, timeout=TIMEOUT)
        self.lock, self.next_ok, self.robots = threading.Lock(), {}, {}
        CACHE.mkdir(parents=True, exist_ok=True)

    def get(self, url: str) -> dict:
        """{'url', 'final', 'status', 'body', 'error'}; a 'problem' label when the request itself failed."""
        path = CACHE / (hashlib.sha1(url.encode()).hexdigest() + ".json.gz")
        if path.exists() and time.time() - path.stat().st_mtime < CACHE_DAYS * 86400:
            return json.loads(gzip.decompress(path.read_bytes()))
        host = urlsplit(url).netloc
        with self.lock:
            wait = max(0.0, self.next_ok.get(host, 0) - time.monotonic())
            self.next_ok[host] = time.monotonic() + wait + DELAY
        time.sleep(wait)
        res = {"url": url, "final": url, "status": 0, "body": "", "error": ""}
        try:
            r = self.client.get(url)
            ctype = r.headers.get("content-type", "")
            res.update(final=str(r.url), status=r.status_code, body=r.text[:3_000_000] if ctype.startswith("text") or not ctype else "")
        except httpx.TimeoutException:
            res["error"] = "the website took too long to answer"
        except httpx.TooManyRedirects:
            res["error"] = "the address redirects in a loop"
        except (httpx.HTTPError, ValueError) as exc:
            msg = str(exc)
            res["error"] = ("its security certificate is broken" if re.search(r"certificate|ssl", msg, re.I) else
                            "the domain does not exist" if re.search(r"name or service|nodename|getaddrinfo|resolve", msg, re.I)
                            else f"could not connect ({msg[:80]})")
        if res["status"] and res["status"] < 500:      # don't remember temporary trouble
            path.write_bytes(gzip.compress(json.dumps(res).encode()))
        return res

    def allowed(self, url: str) -> bool:
        p = urlsplit(url)
        origin = f"{p.scheme}://{p.netloc}"
        if origin not in self.robots:
            r = self.get(origin + "/robots.txt")
            rp = RobotFileParser()
            rp.parse(r["body"].splitlines() if r["status"] == 200 else [])
            self.robots[origin] = rp
        return self.robots[origin].can_fetch(USER_AGENT, url)


@dataclass
class Page:
    url: str
    title: str = ""
    lines: list = field(default_factory=list)
    links: list = field(default_factory=list)       # [(absolute url, link text)]
    jsonld: list = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


def read_page(html: str, url: str) -> Page:
    soup = BeautifulSoup(html, "html.parser")
    page = Page(url, title=ws(soup.title.get_text()) if soup.title else "")
    for s in soup.find_all("script", type=re.compile("ld\\+json", re.I)):
        try:
            page.jsonld.append(json.loads(s.string or ""))
        except ValueError:
            pass
    for t in soup(["script", "style", "noscript", "svg", "template", "iframe", "title"]):
        t.decompose()
    page.links = [(urljoin(url, a["href"].strip()), ws(a.get_text(" "))) for a in soup.find_all("a", href=True)]
    page.lines = [ln for ln in (ws(x) for x in soup.get_text("\n").split("\n")) if ln]
    return page


PARKED = re.compile(r"buy this domain|domain (?:name )?(?:is |may be )?for sale|domain parking|parked (?:free|domain|by)|"
                    r"hugedomains|sedo\.com|afternic|this domain has expired|domain (?:has )?expired|renew (?:this|your) domain|"
                    r"account (?:has been )?suspended|site (?:is )?(?:not|no longer) (?:available|active)", re.I)
SPAM = re.compile(r"(?<![a-z])(?:togel|judi|gacor|slot online|online casino|sportsbook|viagra|cialis|payday loan|porn)(?![a-z])", re.I)
CHURCHY = re.compile(r"(?<![a-z])(?:episcopal|church|parish|worship|chapel|cathedral)", re.I)


def looks_wrong(page: Page) -> str:
    """'' for a church's home page, else why it is not one."""
    text = page.title + "\n" + page.text
    if PARKED.search(text[:6000]):
        return "the address shows a parked, expired or suspended-site page"
    if SPAM.search(text):
        return "the address shows spam or gambling content"
    if len(page.text) >= 300 and not CHURCHY.search(text):
        return "the address shows a website that is not a church's"
    return ""


@dataclass
class Site:
    url: str                          # what was asked for
    pages: list = field(default_factory=list)
    problem: str = ""                 # a PROBLEMS label, or '' when the site was read
    detail: str = ""
    works_at: str = ""                # a working address when the given one is wrong (www., moved domain)


def _alternates(url: str) -> list[str]:
    p = urlsplit(url)
    host = p.netloc[4:] if p.netloc.startswith("www.") else "www." + p.netloc
    other = "http" if p.scheme == "https" else "https"
    return [f"{p.scheme}://{host}{p.path}", f"{other}://{p.netloc}{p.path}"]


def crawl(web: Web, url: str) -> Site:
    site = Site(url)
    shared = any(domain(url) == h or domain(url).endswith("." + h) for h in SHARED_HOSTS)
    scope = (urlsplit(url).path.rstrip("/") or "/") if shared else "/"
    if shared and scope == "/":
        site.problem, site.detail = "Shared site", f"{domain(url)} is a shared (diocese or college) website; this is its home page, not the parish's own page"
        return site
    if not web.allowed(url):
        site.problem, site.detail = "Not readable", "the site's robots.txt asks robots not to visit"
        return site
    r = web.get(url)
    if not r["status"] or r["status"] >= 400:
        for alt in _alternates(url):
            a = web.get(alt)
            if a["status"] == 200 and a["body"]:
                site.works_at, r = alt, a
                break
    if r["error"]:
        site.problem, site.detail = "Can't connect", r["error"]
    elif r["status"] in (401, 403, 429):
        site.problem, site.detail = "Blocked", f"the website refused access (HTTP {r['status']})"
    elif r["status"] in (404, 410):
        site.problem, site.detail = "Not found", f"page not found (HTTP {r['status']})"
    elif r["status"] >= 400:
        site.problem, site.detail = "Can't connect", f"the server answered with an error (HTTP {r['status']})"
    if site.problem:
        return site
    home = read_page(r["body"], r["final"])
    wrong = looks_wrong(home)
    if len(home.text) < 150 and len(home.links) < 5:
        site.problem, site.detail = "Not readable", "the home page has no readable text (it may need JavaScript)"
    elif wrong:
        site.problem, site.detail = "Expired or taken over", wrong
    elif shared and scope != "/" and not urlsplit(r["final"]).path.startswith(scope):
        site.problem, site.detail = "Not found", f"the parish's page on this shared website is gone (it now leads to {r['final']})"
    if site.problem:
        return site
    if domain(r["final"]) != domain(url):
        site.works_at = home_url(r["final"])            # moved to a new domain that still looks like this church
    site.pages.append(home)
    hosts = {urlsplit(url).netloc.lower(), urlsplit(r["final"]).netloc.lower()}
    seen, queue, n = {r["final"].split("#")[0].rstrip("/"), url.rstrip("/")}, [], 0

    def add_links(page: Page, depth: int) -> None:
        nonlocal n
        for link, label in page.links:
            link = link.split("#")[0]
            p = urlsplit(link)
            key = link.rstrip("/")
            if (p.scheme not in ("http", "https") or p.netloc.lower() not in hosts or key in seen or SKIP_URL.search(link)
                    or not p.path.startswith(scope if scope != "/" else "/")):
                continue
            seen.add(key)
            n += 1
            heapq.heappush(queue, (0 if GOOD_PAGE.search(p.path + " " + label) else 1, depth, n, link))

    add_links(home, 1)
    while queue and len(site.pages) < MAX_PAGES:
        _, depth, _, link = heapq.heappop(queue)
        if not web.allowed(link):
            continue
        r = web.get(link)
        if r["status"] == 200 and r["body"] and urlsplit(r["final"]).netloc.lower() in hosts:
            page = read_page(r["body"], r["final"])
            site.pages.append(page)
            if depth < MAX_DEPTH:
                add_links(page, depth + 1)
    return site


# ==================================================================================================
# Facts found on pages
# ==================================================================================================
@dataclass
class Fact:
    kind: str           # phone, email, address, livestream, clergy
    value: str
    url: str
    evidence: str
    role: str = ""      # phone/email: church | rector | ''; clergy: lead | weak | bishop | other
    jsonld: bool = False
    asset_map: bool = False


RECTOR_WORDS = re.compile(r"\b(?:rector|vicar|priest|dean|father|fr|rev|revd|reverend|canon|pastor|clergy|mother|deacon)\b", re.I)
CHURCH_WORDS = re.compile(r"\b(?:office|church|parish|front desk|reception|secretary|administrator|admin|general|info|information)\b", re.I)
WEAK_CHURCH = re.compile(r"\b(?:phone|telephone|tel|call us|call|contact us|main)\b", re.I)


def label_role(text: str, start: int, end: int) -> str:
    """'rector' or 'church' by the NEAREST label word on the same line or the line before; '' if none."""
    ls = text.rfind("\n", 0, start) + 1
    ws_ = max(text.rfind("\n", 0, max(ls - 1, 0)) + 1, start - 160)
    le = text.find("\n", end)
    win = text[ws_: len(text) if le == -1 else le]
    s, e = start - ws_, end - ws_

    def nearest(rx) -> int:
        best = 10 ** 6
        for m in rx.finditer(win):      # words inside the number or address itself ("@gracechurch.org") don't count
            if m.end() <= s:
                best = min(best, s - m.end() + (40 if "\n" in win[m.end():s] else 0))
            elif m.start() >= e:
                best = min(best, m.start() - e + 15)
        return best
    r, c = nearest(RECTOR_WORDS), nearest(CHURCH_WORDS)
    if min(r, c) <= 140:
        return "rector" if r < c else "church"
    return "church" if nearest(WEAK_CHURCH) <= 40 else ""


def email_role(email: str, text: str, start: int, end: int) -> str:
    local = email.split("@")[0]
    if re.search(r"rector|vicar|priest|father|dean|pastor|revd|reverend|\bfr|^rev|^mtr|^mother", local):
        return "rector"
    if re.search(r"office|info|admin|parish|secretary|hello|welcome|contact|church|frontdesk", local):
        return "church"
    return label_role(text, start, end) if start >= 0 else ""


LIVE_HOSTS = ("youtube.com", "facebook.com", "boxcast.tv", "boxcast.com", "resi.io", "vimeo.com", "livestream.com",
              "subsplash.com", "churchstreaming.tv", "online.church")
LIVE_WORDS = re.compile(r"live|stream|watch|worship|service|online|broadcast", re.I)


def livestream_key(url: str) -> str:
    """A channel, whatever tab of it is linked: youtube.com/@x/streams = youtube.com/@x."""
    p = urlsplit(url)
    path = re.sub(r"/(?:live|streams|featured|videos|home)$", "", p.path.rstrip("/"), flags=re.I)
    return (domain(url) + path).lower()


def stable_channel(url: str) -> bool:
    """A link that keeps working: a YouTube channel or a Facebook videos/live page, not one recording."""
    p, d = urlsplit(url), domain(url)
    if d.endswith("youtube.com"):
        return bool(re.match(r"/(?:@|channel/|c/|user/)", p.path))
    if d.endswith("facebook.com"):
        return bool(re.search(r"/(?:live|videos)/?$", p.path))
    return any(d.endswith(h) for h in LIVE_HOSTS)


def page_facts(page: Page) -> list[Fact]:
    text, out = page.text, []
    for m in PHONE_RE.finditer(text):
        before = text[text.rfind("\n", 0, m.start()) + 1: m.start()]
        for d in phone_digits(m.group(0)) if not re.search(r"\bfax\b[^\n]{0,25}$", before, re.I) else ():
            out.append(Fact("phone", d, page.url, ws(text[max(0, m.start() - 80): m.end() + 60]), label_role(text, m.start(), m.end())))
    for m in EMAIL_RE.finditer(text):
        for e in emails_in(m.group(0)):
            out.append(Fact("email", e, page.url, ws(text[max(0, m.start() - 80): m.end() + 40]), email_role(e, text, m.start(), m.end())))
    for href, label in page.links:
        low = href.lower()
        if low.startswith("tel:") and (d := phone_digits(unquote(href[4:]))):
            role = "rector" if RECTOR_WORDS.search(label) else "church" if CHURCH_WORDS.search(label) else ""
            out.append(Fact("phone", d[0], page.url, f"phone link: {label or href}", role))
        elif low.startswith("mailto:"):
            for e in emails_in(unquote(href[7:].split("?")[0])):
                role = email_role(e, "", -1, -1) or ("rector" if RECTOR_WORDS.search(label) else "church" if CHURCH_WORDS.search(label) else "")
                out.append(Fact("email", e, page.url, f"email link: {label or e}", role))
        elif any(domain(href).endswith(h) for h in LIVE_HOSTS) and stable_channel(href):
            out.append(Fact("livestream", href.split("#")[0], page.url, f'link "{label}" -> {href}',
                            "labelled" if LIVE_WORDS.search(label + " " + urlsplit(href).path) else ""))
    for addr, ev in find_addresses(text):
        out.append(Fact("address", addr, page.url, ev))
    for name, role, ev in find_clergy(page.lines):
        kind = "bishop" if re.search(r"\b(?:rt\.?|right|most)\s+rev|\bbishop\b", name, re.I) else clergy_kind(role)
        out.append(Fact("clergy", house_title(name), page.url, ev, kind))
    for d in _walk(page.jsonld):
        if not re.search(r"church|organization|place|localbusiness", str(d.get("@type", "")), re.I):
            continue
        for ph in phone_digits(str(d.get("telephone", ""))):
            out.append(Fact("phone", ph, page.url, f"page data: telephone {d.get('telephone')}", "church", jsonld=True))
        for e in emails_in(str(d.get("email", ""))):
            out.append(Fact("email", e, page.url, f"page data: email {e}", "church", jsonld=True))
        a = d.get("address")
        if isinstance(a, dict) and a.get("streetAddress") and state_code(str(a.get("addressRegion", ""))):
            addr = f"{fmt_street(str(a['streetAddress']))}, {ws(a.get('addressLocality'))}, {state_code(str(a['addressRegion']))}"
            out.append(Fact("address", addr, page.url, f"page data: address {addr}", jsonld=True))
    return out


def _walk(x):
    if isinstance(x, list):
        for y in x:
            yield from _walk(y)
    elif isinstance(x, dict):
        yield x
        for v in x.values():
            if isinstance(v, (dict, list)):
                yield from _walk(v)


def asset_map_facts(html: str, url: str) -> tuple[list[Fact], str]:
    """Facts from an Episcopal Asset Map listing, and the church website it links to ('' if none). Its contact
    person counts as the rector only when the title says Rector, Vicar, Priest-in-Charge, Dean or Interim Rector."""
    soup, out = BeautifulSoup(html, "html.parser"), []

    def txt(sel, root=None):
        el = (root or soup).select_one(sel)
        return ws(el.get_text(" ")) if el else ""
    street = ws(txt(".address-line1") + " " + txt(".address-line2"))
    if street and state_code(txt(".administrative-area")):
        addr = f"{fmt_street(street)}, {txt('.locality')}, {state_code(txt('.administrative-area'))}"
        out.append(Fact("address", addr, url, f"Asset Map address: {addr}", asset_map=True))
    for a in soup.select(".field--name-field-phone a[href^='tel:']")[:1]:
        for d in phone_digits(unquote(a["href"][4:])):
            out.append(Fact("phone", d, url, f"Asset Map phone: {fmt_phone(d)}", "church", asset_map=True))
    for a in soup.select(".field--name-field-email-address a[href^='mailto:']")[:1]:
        for e in emails_in(unquote(a["href"][7:])):
            out.append(Fact("email", e, url, f"Asset Map email: {e}", "church", asset_map=True))
    for p in soup.select(".block-field_contacts .details--person, [data-block-plugin-id*='field_contacts'] .details--person"):
        name, title = txt(".field--name-field-person-name", p), txt(".field--name-field-person-position", p)
        if name and clergy_kind(title) == "lead" and not re.search(r"assistant\s+to|secretary|administrator|director", title, re.I):
            ev = f"Asset Map contact: {name} ({title})"
            out.append(Fact("clergy", house_title(name), url, ev, "lead", asset_map=True))
            for a in p.select("a[href^='mailto:']")[:1]:
                out += [Fact("email", e, url, ev, "rector", asset_map=True) for e in emails_in(unquote(a["href"][7:]))]
            for a in p.select("a[href^='tel:']")[:1]:
                out += [Fact("phone", d, url, ev, "rector", asset_map=True) for d in phone_digits(unquote(a["href"][4:]))]
    link = soup.select_one(".field--name-field-external-url a[href]")
    return out, (link["href"].strip() if link else "")


# ==================================================================================================
# Deciding what to change (only what the website makes certain)
# ==================================================================================================
def _pick(groups: dict, key):
    """The one best (value, facts) by `key`, or None when two values tie for best."""
    ranked = sorted(groups.items(), key=lambda kv: key(kv[1]), reverse=True)
    if not ranked or (len(ranked) > 1 and key(ranked[0][1]) == key(ranked[1][1])):
        return None
    return ranked[0]


def _group(facts: list[Fact], kind: str) -> dict:
    g = defaultdict(list)
    for f in facts:
        if f.kind == kind:
            g[f.value].append(f)
    return g


def sitewide(fs: list[Fact]) -> bool:
    """Shown on 2+ pages (a header or footer), in the page data, or on the Asset Map listing."""
    return len({f.url for f in fs}) >= 2 or any(f.jsonld or f.asset_map for f in fs)


def church_contact(facts: list[Fact], kind: str):
    """The church's main phone/email: labelled for the church on a contact or home page, or shown site-wide."""
    groups = {v: fs for v, fs in _group(facts, kind).items()
              if sum(f.role == "rector" for f in fs) <= sum(f.role == "church" for f in fs)
              and not (kind == "email" and NOT_MAIN_EMAIL.search(v.split("@")[0]))}
    contact = lambda f: f.role == "church" and re.search(r"^/?$|contact|about|visit|office", urlsplit(f.url).path, re.I)
    key = lambda fs: (any(f.role == "church" or f.jsonld for f in fs), len({f.url for f in fs}))
    best = _pick(groups, key)
    return best if best and (sitewide(best[1]) or any(contact(f) for f in best[1])) else None


def rector_contact(facts: list[Fact], kind: str, rector: str):
    """The rector's own phone/email: labelled for the rector AND showing the rector's name (or a rector mailbox)."""
    given, surname = person(rector)
    if not surname:
        return None

    def named(f: Fact) -> bool:
        if f.asset_map:
            return True
        if kind == "email":
            local = f.value.split("@")[0]
            return surname in local or any(len(g) >= 3 and g in local for g in given) or local in ("rector", "vicar", "dean")
        return re.search(r"(?<![a-z])" + re.escape(surname) + r"(?![a-z])", f.evidence.lower()) is not None
    groups = {v: fs for v, fs in _group(facts, kind).items() if any(f.role == "rector" and named(f) for f in fs)}
    return _pick(groups, lambda fs: len({f.url for f in fs}))


def lead_priest(facts: list[Fact]):
    """(name, fact) of the one lead priest the site names on a staff-type page; None if none or several."""
    cl = [f for f in facts if f.kind == "clergy" and not NOT_TODAY.search(urlsplit(f.url).path)]
    leads = [f for f in cl if f.role == "lead"]
    if not leads:   # "Priest" alone counts when it is the only priest named
        weak = [f for f in cl if f.role == "weak"]
        leads = weak if len({person(f.value)[1] for f in weak}) == 1 else []
    leads = [f for f in leads if f.asset_map or STAFF_PAGE.search(urlsplit(f.url).path) or urlsplit(f.url).path in ("", "/")]
    if len({person(f.value)[1] for f in leads}) != 1:
        return None
    return leads[0].value, leads[0]


def _change(rec: dict, fld: str, new: str, f: Fact | None, why: str = "") -> dict:
    cur = rec.get(fld)
    return {"id": rec.get("id"), "name": rec.get("name", ""), "field": fld,
            "action": "add" if blank(cur) else "replace", "current_value": cur, "new_value": new,
            "source_url": f.url if f else "", "evidence": why or (f.evidence[:300] if f else "")}


def decide(rec: dict, facts: list[Fact], site: Site | None) -> list[dict]:
    """The changes for one parish. Values from the church's own website may fill a blank or replace a value the
    site no longer shows anywhere; values from the Asset Map only fill blanks."""
    own = [f for f in facts if not f.asset_map]
    changes = []

    def offer(fld: str, kind: str, choose, fmt, same, avoid=()) -> None:
        cur = rec.get(fld)
        hit = choose(own) or (choose(facts) if blank(cur) else None)
        if not hit or fmt(hit[0]) in avoid:
            return
        if blank(cur) or (sitewide(hit[1]) and not any(same(str(cur), f.value) for f in own if f.kind == kind)):
            changes.append(_change(rec, fld, fmt(hit[0]), hit[1][0]))
    same_phone = lambda cur, v: v in phone_digits(cur)
    same_email = lambda cur, v: v in emails_in(cur)

    # ---- the lead priest first: the rector's phone and email are found by that name ----
    lead = lead_priest(own) or (lead_priest(facts) if blank(rec.get("rector_name")) else None)
    cur = rec.get("rector_name")
    if lead and (blank(cur) or not same_person(str(cur), lead[0])):
        if blank(cur) or not any(f.kind == "clergy" and f.role in ("lead", "weak") and same_person(str(cur), f.value) for f in facts):
            name = lead[0]
            title = re.match(r"(?:The (?:Rt\. |Very )?Rev\.|Fr\.|Mtr\.)(?:\s+Dr\.)?", str(cur or ""))
            if title and person(name)[1] and not re.match(r"(?i)the |fr\.|mtr\.|rev", name):
                name = f"{title.group(0)} {name}"        # "Fr. Raphiell" + "Raphiell Ashford" -> "Fr. Raphiell Ashford"
            changes.append(_change(rec, "rector_name", name, lead[1]))
    rector = lead[0] if lead else ("" if blank(cur) else str(cur))
    offer("church_phone", "phone", lambda fs: church_contact(fs, "phone"), fmt_phone, same_phone)
    offer("church_email", "email", lambda fs: church_contact(fs, "email"), str.lower, same_email)
    if rector:      # never the church's own number or mailbox
        church = {fmt_phone(d) for d in phone_digits(str(rec.get("church_phone") or ""))} | set(emails_in(str(rec.get("church_email") or "")))
        church |= {fmt(h[0]) for fmt, h in ((fmt_phone, church_contact(own, "phone")), (str.lower, church_contact(own, "email"))) if h}
        offer("rector_phone", "phone", lambda fs: rector_contact(fs, "phone", rector), fmt_phone, same_phone, church)
        offer("rector_email", "email", lambda fs: rector_contact(fs, "email", rector), str.lower, same_email, church)

    # ---- address (and its city and state): only when the website shows exactly one street address ----
    addrs = [f for f in own if f.kind == "address"]
    clusters: list[list[Fact]] = []
    for f in addrs:
        home = next((c for c in clusters if addresses_match(c[0].value, f.value)), None)
        home.append(f) if home else clusters.append([f])
    cur = str(rec.get("address") or "")
    num = _street_key(cur)[0]
    gone = not num or not re.search(r"(?<![\w-])" + re.escape(num) + r"(?![\w-])", "\n".join(p.text for p in (site.pages if site else [])), re.I)
    if len(clusters) == 1 and gone and not any(addresses_match(cur, f.value) for f in addrs):
        fs = clusters[0]
        if any(f.jsonld for f in fs) or len({f.url for f in fs}) >= 2:
            common = Counter(f.value for f in fs).most_common(1)[0][0]
            best = next((f for f in fs if f.jsonld), None) or next(f for f in fs if f.value == common)
            changes.append(_change(rec, "address", best.value, best))
            city, st = split_address(best.value)
            if city and city_key(city) != city_key(str(rec.get("city") or "")):
                changes.append(_change(rec, "city", city, best))
            if st and st != state_code(str(rec.get("state") or "")):
                changes.append(_change(rec, "state", st, best))

    # ---- livestream: only to fill a blank, with one channel the site links to as a stream ----
    if blank(rec.get("livestream_url")):
        live = [f for f in own if f.kind == "livestream"]
        streamy = any(re.search(r"live\s?stream|streamed|watch (?:online|live)|watch us", p.text, re.I) for p in (site.pages if site else []))
        live = [f for f in live if f.role == "labelled" or (streamy and domain(f.value).endswith("youtube.com"))]
        groups = defaultdict(list)
        for f in live:
            groups[livestream_key(f.value)].append(f)
        if len(groups) == 1:
            f = next(iter(groups.values()))[0]
            changes.append(_change(rec, "livestream_url", f.value, f))

    # ---- website: the given address moved, or works only with/without www., or is an Asset Map listing ----
    if site and site.works_at:
        changes.append(_change(rec, "website", site.works_at, None,
                               f"{rec.get('website')} does not open, but {site.works_at} does" if domain(site.works_at) == domain(str(rec.get('website')))
                               else f"{rec.get('website')} now leads to {site.works_at}"))
    elif site and domain(site.url) != domain(str(rec.get("website") or "")):
        changes.append(_change(rec, "website", site.url, None, f"the church's own website, linked from its listing {rec.get('website')}"))
    return [c for c in changes if c["new_value"] and c["new_value"] != c["current_value"]]


def bishop_change(rec: dict) -> list[dict]:
    bishop = BISHOPS.get(str(rec.get("diocese") or ""))
    cur = rec.get("diocesan_bishop")
    if bishop and (blank(cur) or not same_person(str(cur), bishop)):
        return [_change(rec, "diocesan_bishop", bishop, None, f"current bishop of the Diocese of {rec['diocese']} (BISHOPS in parish_check.py)")]
    return []


# ==================================================================================================
# The run
# ==================================================================================================
def target_url(rec: dict) -> tuple[str, str]:
    """('own' | 'asset_map' | problem label, url)."""
    url = ws(rec.get("website"))
    if blank(url):
        return "No website", ""
    url = url if "//" in url else "https://" + url
    if domain(url).endswith("facebook.com"):
        return "Not readable", url
    return ("asset_map" if domain(url).endswith("episcopalassetmap.org") else "own"), url


def run(path: Path) -> int:
    records = json.loads(path.read_text(encoding="utf-8-sig"))
    web = Web()
    started = time.monotonic()
    print(f"Checking {len(records)} parishes from {path.name}. Control+C stops; pages read so far are kept for next time.")

    # 1. Asset Map listings: their facts, and the church website each links to
    listings: dict[int, list[Fact]] = {}
    sites_for: dict[int, str] = {}
    problem: dict[int, tuple[str, str, str]] = {}            # record -> (label, details, website)
    for i, rec in enumerate(records):
        kind, url = target_url(rec)
        if kind == "own":
            sites_for[i] = url
        elif kind == "asset_map":
            r = web.get(url) if web.allowed(url) else {"status": 0, "error": "robots.txt asks robots not to visit", "body": ""}
            if r["status"] != 200:
                problem[i] = ("Not found" if r["status"] in (404, 410) else "Can't connect",
                              f"the Asset Map listing could not be read ({r['error'] or 'HTTP ' + str(r['status'])})", url)
                continue
            listings[i], link = asset_map_facts(r["body"], url)
            if link and target_url({"website": link})[0] == "own":
                sites_for[i] = link
        else:
            problem[i] = (kind, "Facebook pages cannot be read without logging in" if url else "", url)

    # 2. Every website once, several at a time
    urls = sorted(set(sites_for.values()))
    sites: dict[str, Site] = {}
    with ThreadPoolExecutor(WORKERS) as pool:
        for n, (u, s) in enumerate(zip(urls, pool.map(lambda u: crawl(web, u), urls)), start=1):
            sites[u] = s
            print(f"\rReading websites {n}/{len(urls)}", end="", flush=True)
    print()

    # 3. Compare
    users = Counter(sites_for.values())
    changes, problems = [], []
    for i, rec in enumerate(records):
        site = sites.get(sites_for.get(i, ""))
        facts = list(listings.get(i, []))
        if site and site.problem:
            problem[i] = (site.problem, site.detail, site.url)
            site = None
        elif site:
            pages = site.pages
            if users[site.url] > 1:      # one website, several parishes: only pages showing this parish's address
                pages = [p for p in pages if any(addresses_match(str(rec.get("address") or ""), a) for a, _ in find_addresses(p.text))]
                if not pages:
                    problem[i] = ("Shared site", f"{users[site.url]} parishes use this website, and no page shows {rec.get('address')}", site.url)
            facts += [f for p in pages for f in page_facts(p)]
            site = Site(site.url, pages, works_at=site.works_at) if pages else None
        changes += decide(rec, facts, site) + bishop_change(rec)
        if i in problem:
            label, details, url = problem[i]
            problems.append({"id": rec.get("id"), "name": rec.get("name", ""), "city": rec.get("city", ""),
                             "website": url, "problem": label, "details": details})

    changes.sort(key=lambda c: (str(c["name"]).casefold(), str(c["id"]), c["field"]))
    problems.sort(key=lambda p: (list(PROBLEMS).index(p["problem"]), str(p["name"]).casefold()))
    write_outputs(path, changes, problems)
    print(f"Done in {round(time.monotonic() - started)} s: {len(changes)} change(s) for {len({c['id'] for c in changes})} parish(es), "
          f"{len(problems)} website problem(s). See {OUT}/")
    return 0


# ==================================================================================================
# Output
# ==================================================================================================
HOW_TO_APPLY = [
    "The file named in input_file is a JSON array of parish records. Each record has a unique \"id\".",
    "For each item in \"changes\": find the record whose \"id\" equals item.id.",
    "Check that record[item.field] equals item.current_value exactly (null means the key is absent). If it does not, "
    "skip the item and report it; the data changed after this check.",
    "Set record[item.field] = item.new_value. action \"add\" fills a blank value; \"replace\" overwrites an outdated one.",
    "Change nothing else: no other fields, no other records, no reordering of keys or records.",
    "Write the file back exactly as Python's json.dumps(records, ensure_ascii=False, indent=2) + \"\\n\" would (UTF-8); "
    "that reproduces its current formatting byte for byte.",
]


def write_outputs(path: Path, changes: list[dict], problems: list[dict]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for old in OUT.iterdir():                     # earlier results (and older versions' files) go; the cache stays
        if old.is_file() and not old.name.startswith((".", "~$")):
            old.unlink()
    try:
        shown = str(path.resolve().relative_to(HERE.parent))
    except ValueError:
        shown = str(path)
    (OUT / "corrections.json").write_text(json.dumps({
        "input_file": shown, "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "how_to_apply": HOW_TO_APPLY, "change_count": len(changes),
        "changes": [{k: c[k] for k in ("id", "field", "action", "current_value", "new_value", "name", "source_url", "evidence")}
                    for c in changes]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    action = {"add": ("Add missing", "D6E6FA"), "replace": ("Replace", "FCE8C3")}
    sheet(OUT / "Corrections.xlsx", "Corrections",
          ["Parish", "Id", "Field", "Change", "Current value", "Correct value", "Source page", "Evidence"],
          [32, 26, 16, 14, 36, 36, 40, 60],
          [([c["name"], c["id"], c["field"], action[c["action"]][0], c["current_value"] or "", c["new_value"],
             c["source_url"], c["evidence"]], action[c["action"]][1]) for c in changes],
          [("Add missing", "D6E6FA", "parishes.json has no value; the website gives one."),
           ("Replace", "FCE8C3", "parishes.json has a value the website no longer shows anywhere; the website gives this one.")])
    sheet(OUT / "Website Problems.xlsx", "Website problems",
          ["Parish", "City", "Id", "Website", "Problem", "Details"], [32, 16, 26, 44, 22, 70],
          [([p["name"], p["city"], p["id"], p["website"], p["problem"], p["details"]], PROBLEMS[p["problem"]][0])
           for p in problems],
          [(k, v[0], v[1]) for k, v in PROBLEMS.items()])


def sheet(path: Path, title: str, headers: list, widths: list, rows: list, key: list) -> None:
    wb = Workbook()
    ws_ = wb.active
    ws_.title = title
    ws_.append(headers)
    for i, w in enumerate(widths):
        ws_.column_dimensions[chr(65 + i)].width = w
        cell = ws_.cell(row=1, column=i + 1)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F3E4E")
    for values, colour in rows:
        ws_.append([("'" + v if isinstance(v, str) and v.startswith("=") else v) for v in values])
        for cell in ws_[ws_.max_row]:
            cell.fill = PatternFill("solid", fgColor=colour)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink, cell.font = cell.value, Font(color="1F5FBF", underline="single")
    ws_.freeze_panes = "A2"
    ws_.auto_filter.ref = ws_.dimensions
    k = wb.create_sheet("Key")
    for label, colour, meaning in key:
        k.append([label, meaning])
        k.cell(row=k.max_row, column=1).fill = PatternFill("solid", fgColor=colour)
        k.cell(row=k.max_row, column=1).font = Font(bold=True)
    k.column_dimensions["A"].width, k.column_dimensions["B"].width = 24, 100
    wb.save(path)


# ==================================================================================================
# Self-test (offline)
# ==================================================================================================
def selftest() -> int:
    fails = []

    def check(name, ok):
        if not ok:
            fails.append(name)
    page = lambda url, text, links=(): Page(url, title="Staff", lines=text.split("\n"), links=list(links))
    check("phone formats", phone_digits("Call (864) 235-5884 or 1.803.556.0100") == ["8642355884", "8035560100"] and fmt_phone("8642355884") == "(864) 235-5884")
    check("address house style", find_addresses("Visit us at 401 North Main Street, Allendale, South Carolina 29810")[0][0] == "401 N Main St, Allendale, SC")
    check("addresses match across spellings", addresses_match("2304 Hwy 17 N, Mt Pleasant, SC", "2304 Highway 17 North, Mount Pleasant, SC")
          and not addresses_match("12 Elm St, A, SC", "14 Elm St, A, SC"))
    check("people", same_person("The Rev. Gary Eichelberger", "Fr. J. Gary Eichelberger") and not same_person("The Rev. J. Smith", "The Rev. M. Smith"))
    check("house title", house_title("REV. JILL WILLIAMS") == "The Rev. Jill Williams" and house_title("The Right Reverend Ann Lee") == "The Rt. Rev. Ann Lee")
    greenwood = "Clergy\nThe Right Reverend Daniel Richards\nBishop of the Diocese of Upper South Carolina\nThe Reverend Jane Rogers Wilson “Janey”\nInterim Rector\nStaff\nDenise Brown\nAdministrative Assistant"
    mtp = "Our Staff\nThe Rt. Rev. Ruth Woodliff-Stanley\nBishop\nThe Rev. Furman Buchanan\nRector\nThe Rev. Brooks Boylan\nAssociate Rector"
    irmo = "LEADERSHIP\nREV. JILL WILLIAMS\nPriest\nSharon Herring\nParish Administrator\nthe REV. JIm neuburger\nAssisting Priest, Retired"
    for text, want in ((greenwood, "The Rev. Jane Rogers Wilson"), (mtp, "The Rev. Furman Buchanan"), (irmo, "The Rev. Jill Williams")):
        lp = lead_priest(page_facts(page("https://church.org/staff", text)))
        check(f"lead priest {want}", lp is not None and lp[0] == want)
    check("vacant post and lay assistant", find_clergy("Vacant\nRector\nLouise Barker\nAssistant to the Rector".split("\n")) == [])
    check("history page clergy ignored", lead_priest(page_facts(page("https://church.org/history", "The Rev. Old Timer\nRector"))) is None)
    rec = {"id": "a", "name": "St. A", "website": "https://church.org/", "address": "12 Elm St, Greenville, SC", "city": "Greenville",
           "state": "SC", "church_phone": "(864) 000-0000", "church_email": "", "rector_name": "The Rev. Ann Lee", "rector_email": "",
           "rector_phone": "", "livestream_url": ""}
    home = page("https://church.org/", "St. A Episcopal Church\n12 Elm Street, Greenville, SC 29601\nOffice: (864) 235-5884\n"
                "Email the office: office@church.org\nWatch our livestream",
                [("https://www.youtube.com/@stachurch", "Watch live"), ("mailto:office@church.org", "office@church.org")])
    staff = page("https://church.org/staff", "The Rev. Ann Lee\nRector\nEmail: alee@church.org\nOffice: (864) 235-5884")
    got = {c["field"]: (c["action"], c["new_value"]) for c in decide(rec, page_facts(home) + page_facts(staff), Site("https://church.org/", [home, staff]))}
    check("phone the site no longer shows is replaced", got.get("church_phone") == ("replace", "(864) 235-5884"))
    check("blank church email filled", got.get("church_email") == ("add", "office@church.org"))
    check("rector email found by the rector's name", got.get("rector_email") == ("add", "alee@church.org"))
    check("livestream channel filled", got.get("livestream_url") == ("add", "https://www.youtube.com/@stachurch"))
    check("matching values are left alone", not {"address", "rector_name", "city"} & set(got))
    got = {c["field"] for c in decide(dict(rec, church_phone="(864) 999-0000"), page_facts(page("https://church.org/", "Call (864) 999-0000\nOffice (864) 235-5884")), None)}
    check("a phone still on the site is never replaced", "church_phone" not in got)
    check("bishop table", bishop_change({"id": "x", "diocese": "South Carolina", "diocesan_bishop": "The Rt. Rev. Old Bishop"})[0]["new_value"] == BISHOPS["South Carolina"])
    got = {c["field"]: c["new_value"] for c in decide(rec, page_facts(page("https://church.org/", "Giving: use our church account give.sta@gmail.com\nCall 555-555-5555\nmymail@mailservice.com")), None)}
    check("no giving mailbox or template placeholders", "church_email" not in got and "church_phone" not in got)
    two = page("https://church.org/contact", "Church: 1644 Highway 174, Edisto Island, SC 29438\nOffice: 42 Station Court, Edisto Island, SC 29438")
    check("highway addresses are read", "1644 Hwy 174, Edisto Island, SC" in [a for a, _ in find_addresses(two.text)])
    off = page("https://church.org/contact", "Church 1004 11th St. Port Royal, SC 29935\nChurch Offices 1110A Paris Ave")
    check("an address still on the site is never replaced",
          "address" not in {c["field"] for c in decide(dict(rec, address="1110A Paris Ave, Port Royal, SC"), page_facts(off) * 2, Site("https://church.org/", [off]))})
    check("rector title kept", decide(dict(rec, rector_name="Fr. Raphiell"), page_facts(page("https://church.org/staff", "Raphiell Ashford, Rector")), None)[0]["new_value"] == "Fr. Raphiell Ashford")
    check("looks wrong", looks_wrong(page("https://x.org/", "This domain is for sale! Buy this domain today.")) != "")
    print("\n".join("FAIL: " + f for f in fails) or "ALL TESTS PASSED")
    return 1 if fails else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Check parishes.json against parish websites (no AI).")
    ap.add_argument("input", nargs="?", help="the parish JSON file to check (required; the live ../data/parishes.json is never assumed)")
    ap.add_argument("--selftest", action="store_true", help="test the rules offline and stop")
    a = ap.parse_args()
    if not a.selftest and not a.input:
        ap.error("name the JSON file to check, e.g.  .venv/bin/python parish_check.py path/to/file.json")
    try:
        sys.exit(selftest() if a.selftest else run(Path(a.input).expanduser()))
    except KeyboardInterrupt:
        print("\nStopped. Pages read so far are cached; run again to continue quickly.")
        sys.exit(130)
