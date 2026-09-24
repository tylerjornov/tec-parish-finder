#!/usr/bin/env python3
"""Convert the parish workbook to data/parishes.json."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Install openpyxl: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = json.loads((ROOT / "data" / "schema.json").read_text())
OUT = ROOT / "data" / "parishes.json"

MAP = SCHEMA["json_fields"]


def slug(name: str, address: str, i: int) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", f"{name}-{address}".lower()).strip("-")
    base = base[:80] or f"parish-{i}"
    return base


def cell(row: dict, header: str):
    v = row.get(header)
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return int(v)
    return str(v).strip()


def num(v):
    if v in ("", None):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main(xlsx_path: str) -> None:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    seen = set()
    for i, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        rec = {headers[j]: raw[j] if j < len(raw) else None for j in range(len(headers))}
        name = cell(rec, "Church Name")
        if not name:
            continue
        addr = cell(rec, "Full Address")
        ident = slug(name, addr, i)
        n = 2
        orig = ident
        while ident in seen:
            ident = f"{orig}-{n}"
            n += 1
        seen.add(ident)
        item = {"id": ident}
        for field, col in MAP.items():
            if field in ("lat", "lon", "asa"):
                item[field] = num(cell(rec, col)) if col in rec else None
            else:
                item[field] = cell(rec, col) if col in rec else ""
        rows.append(item)
    OUT.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n")
    print(f"Wrote {len(rows)} parishes → {OUT}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 scripts/xlsx_to_json.py Episcopal_Churches_US.xlsx")
    main(sys.argv[1])
